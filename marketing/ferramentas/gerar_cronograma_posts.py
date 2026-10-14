# Gera a planilha de produção dos 12 posts de 19/10 a 13/11 (calendário de 12 posts, 3 por semana).
# Conteúdo em conteudo_posts_19-10_a_13-11.py. Uso:
#   python3 marketing/ferramentas/gerar_cronograma_posts.py marketing/Cronograma_Posts_19-10_a_13-11.xlsx
# Depois, recalcular as fórmulas (LibreOffice) para os valores aparecerem em quem lê sem abrir no Excel.
import datetime as dt
import math
import os
import runpy
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = sys.argv[1]
C = runpy.run_path(os.path.join(os.path.dirname(__file__), 'conteudo_posts_19-10_a_13-11.py'))
POSTS, BLOCOS, VALIDACOES = C['POSTS'], C['BLOCOS'], C['VALIDACOES']
D = C['D']

F = 'Arial'
TIF, VERDE, CREME, TCLARO, VLOGO, BRANCO = '96C5BD', '1F3A36', 'F0F0E9', 'BFE0D9', '2F6B64', 'FFFFFF'
fill = lambda c: PatternFill('solid', fgColor=c)
fino = Side(style='thin', color='C9D6D2')
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
BORDA_VERDE = Border(*(Side(style='thin', color=VERDE),) * 4)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
F_NORMAL = Font(name=F, size=10, color=VERDE)
F_CAB = Font(name=F, size=10, bold=True, color=VERDE)
F_TIT = Font(name=F, size=14, bold=True, color=VERDE)
F_SUB = Font(name=F, size=10, italic=True, color=VLOGO)
F_PEQ = Font(name=F, size=8, color=VLOGO)
AMARELO = fill('FFFF00')  # células de parâmetro (convenção da planilha do calendário editorial)
DATA = 'DD/MM/YYYY'
N = len(POSTS)
R0 = 5                      # primeira linha de posts no Cronograma e na aba Posts
R1 = R0 + N - 1
FAIXA = lambda col: f'Cronograma!${col}${R0}:${col}${R1}'


def cab(ws, linha, nomes, larguras=None):
    for j, nome in enumerate(nomes, 1):
        c = ws.cell(row=linha, column=j, value=nome)
        c.font, c.fill, c.border, c.alignment = F_CAB, fill(TIF), BORDA, Alignment(wrap_text=True, vertical='center')
        if larguras:
            ws.column_dimensions[L(j)].width = larguras[j - 1]
    ws.row_dimensions[linha].height = 30


def titulo(ws, texto, sub=None):
    ws['A1'] = texto
    ws['A1'].font = F_TIT
    if sub:
        ws['A2'] = sub
        ws['A2'].font = F_SUB


def altura(valores, larguras, minimo=15):
    # estima a altura da linha a partir do texto mais longo (o Excel não ajusta sozinho linhas geradas por script)
    linhas = 1
    for v, w in zip(valores, larguras):
        if not isinstance(v, str) or v.startswith('='):
            continue
        n = sum(max(1, math.ceil(len(p) / max(1.0, w * 1.15))) for p in v.split('\n'))
        linhas = max(linhas, n)
    return max(minimo, linhas * 13 + 6)


wb = Workbook()

# ---------------------------------------------------------------- Listas
wl = wb.active
wl.title = 'Listas'
titulo(wl, 'Listas e datas de apoio')
FERIADOS = [(D(12, 10), 'Nossa Senhora Aparecida'), (D(2, 11), 'Finados'), (D(15, 11), 'Proclamação da República'), (D(20, 11), 'Consciência Negra')]
ESPECIAIS = [(D(2, 11), 'Finados (feriado)'), (D(14, 11), 'Dia Mundial do Diabetes'), (D(15, 11), 'Proclamação da República (feriado)')]
FORMATOS = ['Reel', 'Carrossel', 'Foto única']
STATUS = ['Ideia', 'Roteiro em revisão', 'Roteiro pronto', 'Aguardando decisão', 'Em produção', 'Arte pronta', 'Em aprovação', 'Aprovado', 'Agendado', 'Publicado']
SITUACAO = ['Aberta', 'Resolvida']
for j, (t, w) in enumerate([('Feriados (fora dos prazos)', 14), ('', 26), ('', 3), ('Datas especiais (no calendário)', 14), ('', 30), ('', 3),
                            ('Formatos', 14), ('Status', 20), ('Blocos', 20), ('Situação', 12)], 1):
    wl.column_dimensions[L(j)].width = w
    if t:
        c = wl.cell(row=2, column=j, value=t)
        c.font, c.fill = F_CAB, fill(TIF)
for i, (d, nome) in enumerate(FERIADOS, 3):
    wl.cell(row=i, column=1, value=d).number_format = DATA
    wl.cell(row=i, column=2, value=nome)
for i, (d, nome) in enumerate(ESPECIAIS, 3):
    wl.cell(row=i, column=4, value=d).number_format = DATA
    wl.cell(row=i, column=5, value=nome)
for col, lista in ((7, FORMATOS), (8, STATUS), (9, [b[0] for b in BLOCOS]), (10, SITUACAO)):
    for i, v in enumerate(lista, 3):
        wl.cell(row=i, column=col, value=v)
for row in wl.iter_rows(min_row=3, max_row=12, max_col=10):
    for c in row:
        c.font = F_NORMAL
FER = f'Listas!$A$3:$A${2 + len(FERIADOS)}'
ESP_D, ESP_N = f'Listas!$D$3:$D${2 + len(ESPECIAIS)}', f'Listas!$E$3:$E${2 + len(ESPECIAIS)}'


def lista_dv(ws, col_lista, n, alvo):
    dv = DataValidation(type='list', formula1=f'=Listas!${col_lista}$3:${col_lista}${2 + n}', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(alvo)


# ---------------------------------------------------------------- Cronograma
wc = wb.create_sheet('Cronograma', 0)
titulo(wc, 'Cronograma de produção e publicação · 19/10 a 13/11 · 12 posts, 3 por semana')
GANTT_INI, NDIAS = 13, 35   # coluna M; 35 dias, de 12/10 a 15/11
larg_c = [6, 11, 6, 8, 18, 11, 44, 11, 11, 11, 18, 14]
cab(wc, 4, ['ID', 'Publicação', 'Dia', 'Semana', 'Bloco', 'Formato', 'Título', 'Roteiro até', 'Criação até', 'Aprovação até', 'Status', 'Responsável'], larg_c)
# parâmetros
params = [('A2', 'Início do mês', 'B2', D(19, 10), DATA, 'Segunda-feira da semana 1. Muda a coluna Semana e o calendário.'),
          ('D2', 'Roteiro: dias úteis antes', 'F2', 5, '0', 'Prazo para o roteiro ficar pronto, em dias úteis antes da publicação.'),
          ('G2', 'Criação (arte ou gravação): dias úteis antes', 'H2', 2, '0', 'Prazo para a arte ou o vídeo ficar pronto.'),
          ('I2', 'Aprovação: dias úteis antes', 'J2', 1, '0', 'Prazo para aprovar a peça final.')]
wc['A2'] = None
for lab, txt, cel, val, fmt, com in params:
    wc[lab] = txt
    wc[lab].font = F_CAB
    wc[lab].alignment = Alignment(wrap_text=True, vertical='center', horizontal='right' if lab != 'A2' else 'left')
    wc[cel] = val
    wc[cel].fill, wc[cel].font, wc[cel].number_format, wc[cel].border = AMARELO, Font(name=F, size=10, bold=True, color='0000FF'), fmt, BORDA
    wc[cel].comment = Comment(com, 'Planejamento')
wc.merge_cells('D2:E2')
wc.row_dimensions[2].height = 42
wc['A3'] = 'Barras do cronograma: R roteiro · C criação (arte ou gravação) · A aprovação · P publicação. Os prazos descontam fins de semana e feriados (aba Listas).'
wc['A3'].font = F_SUB
wc.merge_cells('A3:L3')
# cabeçalho de datas do gráfico
for k in range(NDIAS):
    col = GANTT_INI + k
    wc.column_dimensions[L(col)].width = 3.6
    top = wc.cell(row=3, column=col, value=f'=CHOOSE(WEEKDAY({L(col)}4,2),"S","T","Q","Q","S","S","D")')
    top.font, top.alignment = F_PEQ, CENTRO
    v = f'=$B$2-7' if k == 0 else f'={L(col - 1)}4+1'
    c = wc.cell(row=4, column=col, value=v)
    c.number_format, c.font, c.fill, c.border = 'DD/MM', Font(name=F, size=8, bold=True, color=VERDE), fill(TIF), BORDA
    c.alignment = Alignment(text_rotation=90, horizontal='center', vertical='center')
wc.row_dimensions[4].height = 42
ULT = L(GANTT_INI + NDIAS - 1)
for i, p in enumerate(POSTS):
    r = R0 + i
    vals = [p['id'], p['data'], f'=CHOOSE(WEEKDAY(B{r},2),"seg","ter","qua","qui","sex","sáb","dom")', f'=INT((B{r}-$B$2)/7)+1',
            p['bloco'], p['fmt'], p['titulo'],
            p.get('rot') or f'=WORKDAY(B{r},-$F$2,{FER})', p.get('cri') or f'=WORKDAY(B{r},-$H$2,{FER})', f'=WORKDAY(B{r},-$J$2,{FER})',
            p['status'], '']
    for j, v in enumerate(vals, 1):
        c = wc.cell(row=r, column=j, value=v)
        c.font, c.border, c.alignment = F_NORMAL, BORDA, Alignment(vertical='center', wrap_text=True)
        if j in (2, 8, 9, 10):
            c.number_format = DATA
        if j in (3, 4):
            c.alignment = Alignment(horizontal='center', vertical='center')
    if p.get('rot'):
        for j in (8, 9):
            wc.cell(row=r, column=j).font = Font(name=F, size=10, color='0000FF')
        wc.cell(row=r, column=8).comment = Comment('Data fixa: os dois reels do Dr. Ronald (P01 e P10) têm roteiro neste arquivo e são gravados no mesmo dia.', 'Planejamento')
    for k in range(NDIAS):
        col = L(GANTT_INI + k)
        c = wc.cell(row=r, column=GANTT_INI + k,
                    value=f'=IF({col}$4=$B{r},"P",IF({col}$4=$J{r},"A",IF({col}$4=$I{r},"C",IF({col}$4=$H{r},"R",""))))')
        c.font, c.alignment, c.border = Font(name=F, size=8, bold=True, color=VERDE), CENTRO, BORDA
    wc.row_dimensions[r].height = 30
G = f'M{R0}:{ULT}{R1}'
for letra, cor, fonte in (('P', VERDE, BRANCO), ('A', VLOGO, BRANCO), ('C', TIF, VERDE), ('R', TCLARO, VERDE)):
    wc.conditional_formatting.add(G, CellIsRule(operator='equal', formula=[f'"{letra}"'], fill=fill(cor), font=Font(name=F, size=8, bold=True, color=fonte)))
wc.conditional_formatting.add(f'M3:{ULT}{R1}', FormulaRule(formula=['WEEKDAY(M$4,2)>5'], fill=fill('EEEEEE')))
lista_dv(wc, 'I', len(BLOCOS), f'E{R0}:E{R1}')
lista_dv(wc, 'G', len(FORMATOS), f'F{R0}:F{R1}')
lista_dv(wc, 'H', len(STATUS), f'K{R0}:K{R1}')
wc.freeze_panes = 'H5'
# resumo
r = R1 + 2
wc.cell(row=r, column=1, value='Resumo').font = F_CAB
for k, (rot, f) in enumerate([('Reels', f'=COUNTIF(F{R0}:F{R1},"Reel")'), ('Carrosséis', f'=COUNTIF(F{R0}:F{R1},"Carrossel")'),
                               ('Fotos únicas', f'=COUNTIF(F{R0}:F{R1},"Foto única")'), ('Total', f'=COUNTA(A{R0}:A{R1})'),
                               ('Publicados', f'=COUNTIF(K{R0}:K{R1},"Publicado")')]):
    wc.cell(row=r + 1 + k, column=1, value=rot).font = F_NORMAL
    c = wc.cell(row=r + 1 + k, column=2, value=f)
    c.font = F_CAB
wc.sheet_view.zoomScale = 90
wc.page_setup.orientation = 'landscape'
wc.page_setup.fitToWidth = 1
wc.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Calendário
wk = wb.create_sheet('Calendário', 1)
titulo(wk, 'Calendário do mês · 19/10 a 13/11', 'Os títulos vêm da aba Cronograma: mudou a data lá, muda aqui. A cor mostra o formato.')
wk.column_dimensions['A'].width = 14
for c in range(2, 9):
    wk.column_dimensions[L(c)].width = 22
for j, (t, bg, fg, borda) in enumerate([('Reel', TIF, VERDE, None), ('Carrossel', VERDE, BRANCO, None), ('Foto única', BRANCO, VERDE, BORDA_VERDE)]):
    c = wk.cell(row=3, column=6 + j, value=t)
    c.fill, c.font, c.alignment = fill(bg), Font(name=F, size=10, bold=True, color=fg), CENTRO
    if borda:
        c.border = borda
for j, d in enumerate(['seg', 'ter', 'qua', 'qui', 'sex', 'sáb', 'dom'], 2):
    c = wk.cell(row=4, column=j, value=d)
    c.font, c.fill, c.alignment, c.border = F_CAB, fill(TIF), CENTRO, BORDA
c = wk.cell(row=4, column=1, value='Semana')
c.font, c.fill, c.alignment, c.border = F_CAB, fill(TIF), CENTRO, BORDA
for w in range(4):
    r = 5 + w * 4
    wk.cell(row=r, column=1, value=f'Semana {w + 1}').font = F_CAB
    wk.cell(row=r + 1, column=1, value=f'=TEXT(B{r},"dd/mm")&" a "&TEXT(H{r},"dd/mm")').font = F_SUB
    for c in range(7):
        col = L(2 + c)
        d = wk.cell(row=r, column=2 + c, value=f'=Cronograma!$B$2+{w * 7 + c}')
        d.number_format, d.font, d.alignment = 'DD/MM', F_PEQ, Alignment(horizontal='left', vertical='bottom')
        t = wk.cell(row=r + 1, column=2 + c, value=(
            f'=IFERROR(INDEX({FAIXA("G")},MATCH({col}{r},{FAIXA("B")},0)),IFERROR(INDEX({ESP_N},MATCH({col}{r},{ESP_D},0)),""))'))
        t.font, t.alignment = Font(name=F, size=10, bold=True, color=VERDE), CENTRO
        f2 = wk.cell(row=r + 2, column=2 + c, value=(
            f'=IFERROR(INDEX({FAIXA("F")},MATCH({col}{r},{FAIXA("B")},0))&" · "&INDEX({FAIXA("A")},MATCH({col}{r},{FAIXA("B")},0)),"")'))
        f2.font, f2.alignment = Font(name=F, size=8, color=VERDE), CENTRO
        for rr in (r, r + 1, r + 2):
            wk.cell(row=rr, column=2 + c).border = Border(left=fino, right=fino, top=fino if rr == r else None, bottom=fino if rr == r + 2 else None)
    wk.row_dimensions[r].height = 14
    wk.row_dimensions[r + 1].height = 46
    wk.row_dimensions[r + 2].height = 14
    wk.row_dimensions[r + 3].height = 8
    faixa = f'B{r + 1}:H{r + 2}'
    wk.conditional_formatting.add(faixa, FormulaRule(formula=[f'LEFT(B${r + 2},4)="Reel"'], fill=fill(TIF)))
    wk.conditional_formatting.add(faixa, FormulaRule(formula=[f'LEFT(B${r + 2},4)="Carr"'], fill=fill(VERDE), font=Font(name=F, color=BRANCO, bold=True)))
    lv = Side(style='thin', color=VERDE)
    wk.conditional_formatting.add(f'B{r + 1}:H{r + 1}', FormulaRule(formula=[f'LEFT(B${r + 2},4)="Foto"'], fill=fill(BRANCO), border=Border(left=lv, right=lv, top=lv)))
    wk.conditional_formatting.add(f'B{r + 2}:H{r + 2}', FormulaRule(formula=[f'LEFT(B${r + 2},4)="Foto"'], fill=fill(BRANCO), border=Border(left=lv, right=lv, bottom=lv)))
    # datas especiais (sem post): itálico, sem negrito
    wk.conditional_formatting.add(f'B{r + 1}:H{r + 1}', FormulaRule(formula=[f'AND(B${r + 2}="",B{r + 1}<>"")'], font=Font(name=F, italic=True, bold=False, color=VLOGO)))
r = 22
for k, (rot, f) in enumerate([('Reels', f'=COUNTIF({FAIXA("F")},"Reel")'), ('Carrosséis', f'=COUNTIF({FAIXA("F")},"Carrossel")'),
                               ('Fotos únicas', f'=COUNTIF({FAIXA("F")},"Foto única")')]):
    wk.cell(row=r, column=2 + 2 * k, value=rot).font = F_NORMAL
    c = wk.cell(row=r, column=3 + 2 * k, value=f)
    c.font, c.alignment = F_CAB, Alignment(horizontal='left')
wk.cell(row=23, column=2, value='Ritmo: 1 reel e 2 posts estáticos por semana, segunda, quarta e sexta (na semana de Finados, terça, quarta e sexta).').font = F_SUB
wk.page_setup.orientation = 'landscape'
wk.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Posts (ficha completa)
wp = wb.create_sheet('Posts', 2)
titulo(wp, 'Ficha de cada post', 'Data, dia, semana, bloco, formato, título e status vêm da aba Cronograma. Na arte e na legenda, a marca vai em minúsculas (fairhealth).')
cols_p = [('ID', 6, None), ('Publicação', 11, '=Cronograma!B{r}'), ('Dia', 6, '=Cronograma!C{r}'), ('Semana', 8, '=Cronograma!D{r}'),
          ('Bloco', 16, '=Cronograma!E{r}'), ('Formato', 11, '=Cronograma!F{r}'), ('Título (calendário)', 30, '=Cronograma!G{r}'),
          ('Subtítulo', 32, 'sub'), ('No calendário de 12 posts', 14, 'cal'), ('Especificação', 22, 'espec'),
          ('Gancho / capa (primeira frase)', 28, 'gancho'), ('Objetivo', 30, 'objetivo'), ('O que estamos dizendo por trás', 30, 'portras'),
          ('Persona', 16, 'persona'), ('Fio da história (vem de · leva para)', 34, 'fio'), ('Direção visual', 40, 'visual'),
          ('Material necessário', 32, 'material'), ('Legenda (pronta para colar)', 80, 'legenda'), ('Stories de apoio', 28, 'stories'),
          ('Fonte', 34, 'fonte'), ('Validar antes de publicar', 36, 'validar'), ('Status', 16, '=Cronograma!K{r}'),
          ('Link da arte', 20, ''), ('Comentários da revisão', 28, '')]
cab(wp, 4, [c[0] for c in cols_p], [c[1] for c in cols_p])
for i, p in enumerate(POSTS):
    r = R0 + i
    vals = []
    for nome, w, src in cols_p:
        if nome == 'ID':
            v = p['id']
        elif src is None or src == '':
            v = None
        elif src.startswith('='):
            v = src.format(r=r)
        else:
            v = p[src]
        vals.append(v)
        c = wp.cell(row=r, column=len(vals), value=v)
        c.font, c.border, c.alignment = F_NORMAL, BORDA, WRAP
        if nome == 'Publicação':
            c.number_format = DATA
        if nome in ('Dia', 'Semana'):
            c.alignment = Alignment(horizontal='center', vertical='top')
        if nome in ('Link da arte', 'Comentários da revisão'):
            c.fill = fill(CREME)
    wp.row_dimensions[r].height = altura(vals, [c[1] for c in cols_p])
wp.freeze_panes = 'H5'

# ---------------------------------------------------------------- Roteiros (cena por cena, slide por slide)
wr = wb.create_sheet('Roteiros', 3)
titulo(wr, 'Roteiros: cena por cena (reels) e slide por slide (carrosséis)', 'Título, formato e data vêm da aba Cronograma pelo ID. "Fala" só existe nos reels falados.')
larg_r = [6, 30, 11, 11, 16, 46, 40, 34, 28]
cab(wr, 4, ['ID', 'Post', 'Formato', 'Publicação', 'Parte', 'Texto na tela / no slide', 'Fala (reels)', 'Visual e direção', 'Observação'], larg_r)
r = 5
for i, p in enumerate(POSTS):
    tom = fill(CREME) if i % 2 == 0 else fill(BRANCO)
    for parte, texto, fala, visual, obs in p['roteiro']:
        vals = [p['id'], f'=INDEX({FAIXA("G")},MATCH(A{r},{FAIXA("A")},0))', f'=INDEX({FAIXA("F")},MATCH(A{r},{FAIXA("A")},0))',
                f'=INDEX({FAIXA("B")},MATCH(A{r},{FAIXA("A")},0))', parte, texto, fala or None, visual, obs or None]
        for j, v in enumerate(vals, 1):
            c = wr.cell(row=r, column=j, value=v)
            c.font, c.border, c.alignment, c.fill = F_NORMAL, BORDA, WRAP, tom
            if j == 4:
                c.number_format = DATA
        wr.cell(row=r, column=5).font = F_CAB
        wr.row_dimensions[r].height = altura(vals, larg_r)
        r += 1
wr.freeze_panes = 'F5'

# ---------------------------------------------------------------- Blocos
wbk = wb.create_sheet('Blocos', 4)
titulo(wbk, 'Os 4 blocos da programação', 'Quantidade e datas calculadas a partir da aba Cronograma.')
larg_b = [20, 36, 60, 18, 10, 12, 12]
cab(wbk, 4, ['Bloco', 'Pergunta que responde', 'Intuito', 'Posts', 'Nº de posts', 'Primeira data', 'Última data'], larg_b)
for i, (b, perg, intu, ids) in enumerate(BLOCOS):
    r = 5 + i
    vals = [b, perg, intu, ids, f'=COUNTIF({FAIXA("E")},A{r})', f'=_xlfn.MINIFS({FAIXA("B")},{FAIXA("E")},A{r})', f'=_xlfn.MAXIFS({FAIXA("B")},{FAIXA("E")},A{r})']
    for j, v in enumerate(vals, 1):
        c = wbk.cell(row=r, column=j, value=v)
        c.font, c.border, c.alignment = F_NORMAL, BORDA, WRAP
        if j in (6, 7):
            c.number_format = DATA
    wbk.cell(row=r, column=1).font = F_CAB
    wbk.row_dimensions[r].height = altura(vals, larg_b)
wbk.cell(row=10, column=1, value='Já publicados antes desta programação: "O que é a fairhealth?" e "Conheça o Escritório do Cuidado". O post do time (03) segue segurado.').font = F_SUB

# ---------------------------------------------------------------- Validações
wv = wb.create_sheet('Validações', 5)
titulo(wv, 'O que validar antes de publicar', 'Mude a situação para "Resolvida" quando a validação for feita.')
larg_v = [5, 12, 70, 28, 12, 32]
cab(wv, 4, ['#', 'Post', 'O que validar', 'Quem confirma', 'Situação', 'Observação'], larg_v)
for i, (ids, oque, quem) in enumerate(VALIDACOES):
    r = 5 + i
    vals = [i + 1, ids, oque, quem, 'Aberta', None]
    for j, v in enumerate(vals, 1):
        c = wv.cell(row=r, column=j, value=v)
        c.font, c.border, c.alignment = F_NORMAL, BORDA, WRAP
    wv.cell(row=r, column=6).fill = fill(CREME)
    wv.row_dimensions[r].height = altura(vals, larg_v)
r_fim = 4 + len(VALIDACOES)
lista_dv(wv, 'J', len(SITUACAO), f'E5:E{r_fim}')
wv.conditional_formatting.add(f'E5:E{r_fim}', CellIsRule(operator='equal', formula=['"Resolvida"'], fill=fill(TIF)))
wv.cell(row=r_fim + 2, column=2, value='Abertas').font = F_NORMAL
wv.cell(row=r_fim + 2, column=3, value=f'=COUNTIF(E5:E{r_fim},"Aberta")').font = F_CAB

# ---------------------------------------------------------------- Como usar
wu = wb.create_sheet('Como usar', 0)
wu.column_dimensions['A'].width = 130
linhas = [
    ('Programação do Instagram da fairhealth · 19/10 a 13/11', F_TIT),
    ('12 posts em 4 semanas, 3 por semana (1 reel e 2 posts estáticos). Versão de 14/10/2026. A mesma programação está em PowerPoint: marketing/Programacao_Instagram_19-10_a_13-11_v4.pptx.', F_NORMAL),
    ('', F_NORMAL),
    ('Abas', F_CAB),
    ('Cronograma: a lista dos 12 posts com data, formato e status, os prazos de roteiro, criação e aprovação e um gráfico de barras dia a dia. É a aba principal: as outras puxam os dados dela.', F_NORMAL),
    ('Calendário: o mês em formato de calendário, com uma cor por formato (Reel em Tiffany, Carrossel em verde, Foto única em branco).', F_NORMAL),
    ('Posts: a ficha completa de cada post: subtítulo, gancho, objetivo, o que estamos dizendo por trás, fio da história, direção visual, material, legenda pronta, stories de apoio, fonte e validação.', F_NORMAL),
    ('Roteiros: cada reel cena por cena (tempo, texto na tela, fala e visual) e cada carrossel slide por slide.', F_NORMAL),
    ('Blocos: os 4 blocos (Reconhecimento, Relevância, Compreensão e Confiança), com a pergunta e o intuito de cada um.', F_NORMAL),
    ('Validações: o que precisa de aval antes de publicar, com quem confirma e a situação.', F_NORMAL),
    ('Listas: feriados, datas especiais e os valores das listas suspensas.', F_NORMAL),
    ('', F_NORMAL),
    ('Como editar', F_CAB),
    ('Células amarelas (aba Cronograma, linha 2): a data de início do mês e os prazos de roteiro, criação e aprovação, em dias úteis antes da publicação. Mude ali e todos os prazos se ajustam.', F_NORMAL),
    ('Para mudar a data de um post, mude a coluna Publicação no Cronograma: o calendário, a ficha e os roteiros acompanham.', F_NORMAL),
    ('Status, Formato e Bloco têm lista suspensa. Datas em azul no Cronograma são fixas (os dois reels do Dr. Ronald, gravados no mesmo dia).', F_NORMAL),
    ('Fluxo de status: Ideia → Roteiro em revisão → Roteiro pronto → Em produção → Arte pronta → Em aprovação → Aprovado → Agendado → Publicado (ou Aguardando decisão).', F_NORMAL),
    ('', F_NORMAL),
    ('Regras de texto e arte', F_CAB),
    ('Marca em minúsculas (fairhealth) na arte e na legenda. Só Raleway. O Tiffany é a cor principal e entra como fundo, não só como detalhe. Fio contínuo como elemento gráfico. Sem amarelo nas artes.', F_NORMAL),
    ('Números sempre com fonte e com suavidade (o cuidado vem primeiro). Sem promessa clínica, sem preço e sem prometer conformidade com a NR-1. Nenhum nome de cliente; personagem fictícia sempre sinalizada.', F_NORMAL),
    ('Legendas com 5 hashtags, a primeira sempre #fairhealth.', F_NORMAL),
    ('', F_NORMAL),
    ('Fontes do conteúdo', F_CAB),
    ('marketing/calendario-12-posts.md · instagram/essencia.md · instagram/voice.md · instagram/posts/nossa-historia/README.md · planilha do calendário editorial (roteiros de 09/10 a 13/10).', F_NORMAL),
]
for i, (t, f) in enumerate(linhas, 1):
    c = wu.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, vertical='top')

wb.move_sheet('Listas', offset=len(wb.sheetnames))
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ('Listas',)
wb.active = 1
wb.save(SAIDA)
print('ok', SAIDA)
