# Gera marketing/Calendario_Editorial_FairHealth.xlsx (planilha de trabalho do time).
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = 'Arial'
TIFFANY, TEXTO, AMARELO = '9FC5BD', '1F3A36', 'FFFF00'
thin = Side(style='thin', color='C9D6D2')
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CAB = Font(name=F, bold=True, color=TEXTO, size=10)
NORMAL = Font(name=F, size=10)
TITULO = Font(name=F, bold=True, size=14, color=TEXTO)
FILL_CAB = PatternFill('solid', fgColor=TIFFANY)
FILL_INPUT = PatternFill('solid', fgColor=AMARELO)

I, P, J = 'Institucional', 'Produto', 'Projetos, eventos e entregas'
V, C, FU = 'Vídeo', 'Carrossel', 'Foto única'
IGLI, IG = 'Instagram + LinkedIn', 'Instagram'
RH, COL, SESMT, CFO = 'RH e liderança', 'Colaboradores', 'SESMT e medicina do trabalho', 'Diretoria financeira'
FJ = 'FairJob'
HT = '#fairhealth #saudecorporativa #cuidadocentradonapessoa #medicinadoestilodevida #saudedotrabalhador'

# deslocamento (dias a partir da segunda) das 4 publicações de cada semana; evita feriados nacionais
OFF = {w: (0, 1, 3, 4) for w in range(1, 13)}
OFF[3] = (1, 2, 3, 4)    # 02/11 Finados
OFF[5] = (0, 1, 2, 3)    # 20/11 Consciência Negra
OFF[10] = (0, 1, 2, 3)   # 25/12 Natal
OFF[11] = (0, 1, 2, 3)   # 01/01 Confraternização

def time(n, nome, fmt, sem):
    return dict(
        linha=I, tema='Time', serie=f'Quem cuida ({n}/8)', titulo=f'Quem cuida: {nome}',
        publico=f'{COL}; {RH}',
        nec='Quem vai me atender? Quem são as pessoas por trás da Fair Health?',
        obj='Humanizar a marca e criar vínculo', formato=fmt, canal=IGLI,
        collab=f'{FJ}; perfil de {nome} (se houver e se a pessoa quiser)',
        msg=f'[Resposta de {nome} a "O que é importante para você?"]',
        valid='Material da pessoa: ' + ('vídeo curto (20 a 30 s)' if fmt == V else 'retrato vertical') +
              ', descrição curta sem cargo e resposta à pergunta. Aprovação da pessoa antes de publicar.',
        status='Aguardando material' if sem <= 4 else 'Pauta proposta',
        criativo='Modelo M4 (retrato do time)' if fmt != V else 'Modelo M9 (vídeo, capa com nome)')

ROWS = [
 # ---------------- MÊS 1 ----------------
 dict(mes=1, sem=1, linha=I, tema='Apresentação da Fair Health', serie='',
  titulo='Antes de qualquer exame, uma conversa', publico=f'{COL}; {RH}',
  nec='O que é a Fair Health e o que ela faz dentro da empresa?',
  obj='Apresentar a marca e o jeito de cuidar', formato=V, canal=IGLI, collab=FJ,
  msg='A Fair Health leva o cuidado para dentro da empresa, e todo cuidado começa com uma pergunta: o que é importante para você?',
  valid='Roteiro e legenda citam assistentes sociais: confirmar se fazem parte do atendimento antes de publicar (o Executive Summary cita médicos e psicólogos).',
  status='Roteiro pronto',
  texto='Antes de qualquer exame, uma conversa.\n\nEsse é o Escritório do Cuidado: um espaço dentro da empresa, com médicos, psicólogos e assistentes sociais, perto de quem trabalha.\n\nAqui, todo cuidado começa com a mesma pergunta: o que é importante para você? A partir da resposta, montamos juntos um plano que cabe na sua rotina e seguimos acompanhando.\n\nSeja bem-vindo. Nos próximos dias, você vai conhecer quem cuida por aqui.\n\n' + HT,
  roteiro='Reel de 30 a 34 s (instagram/primeiros-posts-acolhimento.md, ID 1).\n0–3 s: "Antes de qualquer exame, a gente começa com uma conversa." Porta do Escritório se abrindo.\n3–9 s: o espaço dentro da empresa (recepção, copa, sala).\n9–16 s: quem chega é recebido por gente, sem pressa e sem julgamento.\n16–23 s: cada integrante diz um pedaço da pergunta "O que é importante para você?".\n23–30 s: a partir da resposta, um cuidado que cabe na rotina.\n30–34 s: logo e assinatura.',
  criativo='Vídeo 1080x1920 gravado no espaço real e com o time. Texto na tela dentro da área segura. Capa no modelo M1 (Tiffany, pergunta), conferindo o recorte da grade. LinkedIn: mesmo vídeo, legenda adaptada para RH.'),
 dict(mes=1, sem=1, linha=P, tema='Escritório do Cuidado', serie='Como funciona o cuidado (1/3): o percurso',
  titulo='Aqui, o cuidado começa com uma pergunta', publico=COL,
  nec='O que acontece quando eu vou ao Escritório do Cuidado? Vou ser julgado? O que eu contar fica entre nós?',
  obj='Explicar o percurso do atendimento e reduzir a barreira de ir', formato=C, canal=IGLI, collab=FJ,
  msg='Primeiro a gente te ouve e entende sua rotina; depois monta o plano com você e acompanha. O que você conta fica entre você e quem cuida.',
  valid='Responsável médica confere as etapas citadas (anamnese, bioimpedância, smartband por 48 horas, retorno).',
  status='Em revisão',
  texto='Aqui, o cuidado começa com uma pergunta: o que é importante para você?\n\nAntes de exame, antes de número, a gente quer saber da sua rotina, do seu trabalho, do que anda pesando e do que você gostaria de mudar.\n\nO Escritório do Cuidado fica dentro da empresa, perto de onde você já está. Lá, médicos e psicólogos se revezam para montar com você um cuidado que caiba na sua vida, com base na Medicina do Estilo de Vida.\n\nO objetivo da fairhealth é melhorar a qualidade de vida de quem trabalha, com um cuidado feito para cada pessoa e acompanhado de perto.\n\nE o que você conta fica entre você e quem cuida.\n\nArrasta para o lado para ver como funciona.\n\n' + HT,
  roteiro='10 slides: 1 capa (pergunta) · 2 fica dentro da empresa · 3 gente que cuida de gente · 4 Passo 1: a gente te conhece · 5 Passo 2: o que é importante para você? · 6 Passo 3: entender antes de propor (bioimpedância e smartband por 48 h) · 7 Passo 4: o plano é seu · 8 Passo 5: retorno e acompanhamento · 9 sigilo · 10 convite para comentar.',
  criativo='Arte no Canva (design DAHW6JfDldU, 11 páginas): escolher entre as versões originais e as novas dos slides 1, 5, 7, 9 e 10. Modelos M1, M2 e M3. LinkedIn: exportar em PDF.'),
 dict(mes=1, sem=1, linha=J, tema='Escritório em funcionamento', serie='',
  titulo='Um dia no Escritório do Cuidado', publico=f'{RH}; {COL}',
  nec='Como isso funciona na prática, dentro de uma empresa?',
  obj='Mostrar o serviço em operação e dar concretude à oferta', formato=V, canal=IGLI,
  collab=f'{FJ}; empresa-cliente (só se autorizar)',
  msg='O cuidado acontece dentro da empresa, em plantões de médicos e psicólogos, perto de quem trabalha.',
  valid='Gravar no espaço real (as imagens atuais parecem renders). Autorização da empresa para gravar. Nenhum colaborador atendido em cena.',
  status='Aguardando material',
  texto='[Rascunho] Como é um dia no Escritório do Cuidado?\n\nAntes do primeiro atendimento, a sala é preparada, a agenda é revisada e as smartbands ficam prontas para quem vai começar o acompanhamento.\n\nDepois, é conversa. Médicos e psicólogos se revezam em plantões dentro da empresa, perto de quem trabalha.\n\n' + HT,
  roteiro='Reel de 30 a 40 s, sem pacientes.\n0–3 s: porta se abrindo. Tela: UM DIA NO ESCRITÓRIO DO CUIDADO\n3–10 s: profissional preparando a sala (cena C3 da sessão de fotos).\n10–18 s: agenda e smartbands sobre a mesa (A7).\n18–28 s: profissional recebendo alguém, de costas. Tela: PLANTÕES DE MÉDICOS E PSICÓLOGOS\n28–35 s: copa e pausa (B4). Fechamento com a pergunta e o logo.',
  criativo='Vídeo 1080x1920. Cenas B1, B2, B4, C3 e A7 de instagram/sessao-de-fotos.md. Capa no modelo M7.'),
 dict(mes=1, sem=1, **{**time(1, 'Marília', C, 1),
  'texto': 'Por trás de todo cuidado, tem gente. Esta é a Marília.\n\n[Descrição curta, escrita com ela, sem cargo.]\n\nPerguntamos a ela o mesmo que perguntamos a quem cuidamos: o que é importante para você?\n\n"[resposta]"\n\nE para você, o que é importante?\n\n' + HT,
  'roteiro': 'Carrossel de 3 slides: 1 retrato em tela cheia e nome · 2 quem é a Marília (sem cargo) · 3 "O que é importante para você?" e a resposta dela.'}),
 dict(mes=1, sem=2, linha=P, tema='Medicina do Estilo de Vida', serie='Pilares da MEV (abertura)',
  titulo='Medicina do Estilo de Vida em 1 minuto: os 7 pilares', publico=f'{COL}; {RH}',
  nec='O que é Medicina do Estilo de Vida? É dieta? É academia?',
  obj='Apresentar a base do cuidado e abrir a série dos pilares', formato=V, canal=IGLI, collab=FJ,
  msg='Os pilares não funcionam isoladamente: o plano olha para a pessoa inteira.',
  valid='Responsável médica valida a definição de MEV e os nomes dos 7 pilares (cartilha do Escritório do Cuidado).',
  status='Pauta proposta', criativo='Modelo M9 (profissional para a câmera) com etiqueta "PILARES".'),
 dict(mes=1, sem=2, linha=J, tema='Letramento e workshops', serie='',
  titulo='Cuidado também se aprende: como funcionam os workshops nos pilares', publico=f'{RH}; {COL}',
  nec='Além da consulta, o que acontece com o time ao longo do programa?',
  obj='Mostrar uma entrega do programa além do atendimento individual', formato=C, canal=IGLI, collab=FJ,
  msg='Ao longo do programa, o cuidado também acontece em grupo, com letramento e workshops nos pilares da MEV.',
  valid='Time confirma formato, frequência e temas dos workshops. Fotos reais de um encontro, com autorização.',
  status='Aguardando material', criativo='Modelo M7 (bastidor) com fotos reais.'),
 dict(mes=1, sem=2, **time(2, 'Nina', V, 2)),
 dict(mes=1, sem=2, linha=P, tema='Plano Individualizado de Cuidado', serie='Como funciona o cuidado (2/3): o plano',
  titulo='O plano é seu: como nasce o Plano Individualizado de Cuidado', publico=COL,
  nec='Vou receber uma lista de regras impossíveis de seguir?',
  obj='Explicar como o plano é montado com a pessoa (decisão compartilhada)', formato=C, canal=IGLI, collab=FJ,
  msg='O plano parte do que é importante para você e da sua rotina real, com ações nos pilares da MEV, e muda quando a vida muda.',
  valid='Responsável médica valida etapas e prazos de retorno e acompanhamento. Se usar o depoimento A1, confirmar o formato (CFM).',
  status='Pauta proposta', criativo='Modelo M3 (card de interface do plano), já testado no carrossel de acolhimento.'),
 dict(mes=1, sem=3, linha=J, tema='Entregas para o RH', serie='Tecnologia a favor da pessoa',
  titulo='Bastidores: como nasce o relatório que o RH recebe', publico=RH,
  nec='O que eu recebo para acompanhar o programa? Vou ver dados individuais?',
  obj='Mostrar a entrega para a empresa e a regra de privacidade', formato=V, canal=IGLI, collab=FJ,
  msg='O RH acompanha o programa por indicadores do grupo, anonimizados e integrados ao fairdata. Nenhum dado individual.',
  valid='Time confirma quais indicadores entram no relatório. Tela genérica, sem dado real.',
  status='Pauta proposta', criativo='Vídeo narrado pela Nina sobre painel genérico (cena C4). Modelo M9.'),
 dict(mes=1, sem=3, linha=I, tema='Origem', serie='',
  titulo='Tudo começou num hospital: do Escritório do Paciente ao Escritório do Cuidado', publico=f'{RH}; {SESMT}',
  nec='De onde vem esse modelo? É só mais um programa de bem-estar?',
  obj='Dar lastro à marca com a história do modelo', formato=C, canal=IGLI, collab=FJ,
  msg='O modelo nasceu em 2009 para colocar o paciente no centro do cuidado. Hoje, coloca o colaborador, dentro do ecossistema FairJob.',
  valid='[PENDENTE] Quem idealizou o Escritório do Paciente. Conferir as datas da linha do tempo (Aims do IHI) com fonte.',
  status='Pauta proposta', criativo='Modelo M2 (linha do tempo em passos).'),
 dict(mes=1, sem=3, linha=P, tema='Pilar: sono', serie='Pilares da MEV (1: sono)',
  titulo='Dormir não é tempo perdido', publico=COL,
  nec='Por que eu acordo cansado? O que dá para mudar na rotina?',
  obj='Educar sobre o pilar sono e mostrar como ele entra no cuidado', formato=V, canal=IG, collab=FJ,
  msg='Dormir não é tempo perdido, é investimento em saúde. No Escritório do Cuidado, o sono entra no plano a partir da sua rotina real.',
  valid='Responsável médica valida as orientações. Fonte: guia "Saúde mental: o cuidado no dia a dia" (higiene do sono).',
  status='Pauta proposta', criativo='Modelo M9 com etiqueta "PILARES · SONO". Cenas D3 da sessão de fotos.'),
 dict(mes=1, sem=3, linha=J, tema='Como começa numa empresa', serie='Como chega a uma empresa (1/2)',
  titulo='Como o Escritório do Cuidado chega a uma empresa: diagnóstico, proposta e implantação', publico=f'{RH}; {CFO}',
  nec='Quero isso na minha empresa. Por onde começa e o que acontece em cada etapa?',
  obj='Deixar clara a oferta e o primeiro passo comercial', formato=C, canal=IGLI, collab=FJ,
  msg='Começa com um diagnóstico gratuito, segue com uma proposta personalizada e chega à implantação, com onboarding e primeiros atendimentos. Comente CUIDADO.',
  valid='Não citar prazos de cada etapa sem confirmação. Confirmar o destino do link (www.fairhealth.com.br).',
  status='Pauta proposta', criativo='Modelo M2 (3 passos) e chamada CUIDADO no último slide.'),
 dict(mes=1, sem=4, **time(3, 'Ronald', V, 4)),
 dict(mes=1, sem=4, linha=P, tema='Tecnologia: wearables', serie='Como funciona o cuidado (3/3): o acompanhamento',
  titulo='O cuidado não para quando a consulta acaba', publico=f'{COL}; {RH}',
  nec='Para que serve a pulseira? Vão me vigiar?',
  obj='Explicar o acompanhamento com tecnologia e a regra de privacidade', formato=C, canal=IGLI, collab=FJ,
  msg='A smartband e a bioimpedância ajudam a olhar, junto com a pessoa, como foram as últimas semanas. Os dados ficam no atendimento; o RH vê só o grupo.',
  valid='[PENDENTE] Quais dados a smartband registra e se todos os atendidos usam a pulseira (instagram/posts/wearables/README.md).',
  status='Em revisão', criativo='Arte pronta em instagram/posts/wearables (8 slides); a capa ainda espera a foto A1. Adaptar ao visual atual.'),
 dict(mes=1, sem=4, linha=J, tema='Registro do mês', serie='Registro do mês',
  titulo='[Registro do mês: workshop, evento ou entrega, a confirmar]', publico='A definir',
  nec='O que a Fair Health está fazendo agora?', obj='Mostrar movimento real e prova social',
  formato=V, canal=IGLI, collab=f'{FJ}; parceiros envolvidos (se autorizarem)', msg='A definir',
  valid='Agenda de eventos e entregas do mês. Autorizações de imagem.', status='Aguardando pauta',
  criativo='Modelo M7 (bastidor).'),
 dict(mes=1, sem=4, **time(4, 'Floriana', C, 4)),

 # ---------------- MÊS 2 ----------------
 dict(mes=2, sem=5, linha=J, tema='Implantação', serie='Como chega a uma empresa (2/2)',
  titulo='Bastidores de uma implantação: os primeiros atendimentos numa empresa', publico=RH,
  nec='Como é, na prática, o começo do Escritório do Cuidado numa empresa?',
  obj='Mostrar uma implantação real, continuando o post M1-12', formato=V, canal=IGLI,
  collab=f'{FJ}; empresa-cliente (só se autorizar)', msg='Do onboarding aos primeiros atendimentos, o cuidado chega perto de quem trabalha.',
  valid='Empresa, datas e autorização de imagem a confirmar. Nenhum colaborador atendido identificável.',
  status='Aguardando pauta', criativo='Modelo M7 (bastidor).'),
 dict(mes=2, sem=5, linha=P, tema='NR-1 e riscos psicossociais', serie='NR-1 (1/4): contexto',
  titulo='NR-1: o que muda quando os riscos psicossociais entram na gestão', publico=f'{RH}; {SESMT}',
  nec='O que a NR-1 passa a pedir da minha empresa e por onde eu começo?',
  obj='Informar com rigor e mostrar que cumprir a norma é mudar o dia a dia, não só preencher documento', formato=C, canal=IGLI, collab=FJ,
  msg='Adequar-se à NR-1 não é só preencher a planilha de riscos. Se nada muda no dia a dia de quem trabalha, a empresa continua exposta.',
  valid='Texto da norma, prazos e obrigações com fonte oficial (Ministério do Trabalho). A mensagem usa a posição [SUGESTÃO] 1 do voice.md, que precisa de aprovação. Revisão da responsável médica.',
  status='Pauta proposta', criativo='Modelo M2 (passos) com fonte no rodapé (M6).'),
 dict(mes=2, sem=5, **time(5, 'Renata', V, 5)),
 dict(mes=2, sem=5, linha=P, tema='Pilar: saúde mental e estresse', serie='Pilares da MEV (2: saúde mental e estresse)',
  titulo='Quando o estresse não passa no fim de semana', publico=COL,
  nec='Isso que eu sinto é normal? Quando vale procurar ajuda?',
  obj='Educar com fonte, acolher e mostrar que há psicólogos por perto', formato=C, canal=IGLI, collab=FJ,
  msg='Estresse que não passa merece atenção. No Escritório do Cuidado, médicos e psicólogos estão perto para conversar. Em crise, o CVV atende pelo 188.',
  valid='Responsável médica. Fontes do guia (OMS, APA, IPq-USP). Fechar com onde pedir ajuda (CVV 188, CAPS).',
  status='Pauta proposta', criativo='Modelo M5 (pilar).'),
 dict(mes=2, sem=6, **time(6, 'Liliane', V, 6)),
 dict(mes=2, sem=6, linha=J, tema='Resultados liberados', serie='Da intenção à evidência (1: adesão e retorno)',
  titulo='Chemitec, 1º semestre de 2026: as pessoas usam e voltam', publico=f'{RH}; {CFO}',
  nec='Os colaboradores usam de verdade um programa como esse?',
  obj='Provar adesão com dado real e agregado', formato=C, canal=IGLI, collab=f'{FJ}; Chemitec (se autorizar a marcação)',
  msg='161 atendimentos e 48 colaboradores atendidos (98% de adesão); 81,3% voltaram para dois ou mais atendimentos.',
  valid='Dados liberados (voice.md), sempre agregados. Confirmar se a empresa autoriza ser marcada ou convidada para a collab.',
  status='Pauta proposta', criativo='Modelo M6 (dado com fonte).'),
 dict(mes=2, sem=6, linha=P, tema='NR-1 para quem trabalha', serie='NR-1 (versão para colaboradores)',
  titulo='Seu bem-estar no trabalho também é assunto da empresa', publico=COL,
  nec='O que essa tal de NR-1 tem a ver comigo?',
  obj='Traduzir a NR-1 para quem trabalha e apresentar o Escritório como canal de cuidado', formato=V, canal=IG, collab=FJ,
  msg='A NR-1 pede que a empresa olhe também para o que pesa no dia a dia de quem trabalha. O Escritório do Cuidado é um dos lugares onde isso acontece.',
  valid='Texto da norma com fonte oficial. Revisão da responsável médica.', status='Pauta proposta',
  criativo='Modelo M9 (profissional para a câmera).'),
 dict(mes=2, sem=6, linha=I, tema='Ecossistema FairJob', serie='',
  titulo='Fair Health no ecossistema FairJob: medir, cuidar, provar, propagar', publico=RH,
  nec='Quem está por trás da Fair Health e como ela se conecta ao resto?',
  obj='Situar a marca no ecossistema e reforçar a parceria com a FairJob', formato=C, canal=IGLI, collab=FJ,
  msg='Dentro do ecossistema FairJob, a Fair Health é o cuidar.',
  valid='Confirmar com a FairJob como o ecossistema é descrito hoje e quais marcas citar.', status='Pauta proposta',
  criativo='Modelo M2.'),
 dict(mes=2, sem=7, linha=P, tema='NR-1 e riscos psicossociais', serie='NR-1 (2/4): diagnóstico',
  titulo='Como enxergar o que pesa no dia a dia do time', publico=f'{RH}; {SESMT}',
  nec='Como identifico os riscos psicossociais sem expor ninguém?',
  obj='Mostrar o papel da escuta e dos indicadores anonimizados no diagnóstico', formato=V, canal=IGLI, collab=FJ,
  msg='Escuta individual e indicadores do grupo, anonimizados, mostram onde agir sem expor ninguém.',
  valid='Responsável médica. Confirmar o que a plataforma de indicadores entrega (Executive Summary cita programa de prevenção de riscos psicossociais e plataforma de indicadores).',
  status='Pauta proposta', criativo='Modelo M9.'),
 dict(mes=2, sem=7, linha=J, tema='Depoimentos', serie='Da intenção à evidência (2: experiência de quem é cuidado)',
  titulo='"O processo não precisa ser perfeito para estar funcionando."', publico=f'{COL}; {RH}',
  nec='Como é ser cuidado pelo Escritório do Cuidado?',
  obj='Mostrar a experiência de quem é atendido, sem prometer resultado', formato=FU, canal=IGLI, collab=FJ,
  msg='Cuidado centrado na pessoa: constância vale mais que perfeição.',
  valid='Trecho B1 aprovado (instagram/depoimentos.md). Confirmar o formato com a responsável médica (publicidade médica, CFM). Assinatura padrão e logo em coração.',
  status='Pauta proposta', criativo='Modelo M8 (citação).'),
 dict(mes=2, sem=7, **time(7, 'Fernando', V, 7)),
 dict(mes=2, sem=7, linha=P, tema='Tecnologia e dados', serie='Tecnologia a favor da pessoa',
  titulo='Seus dados, seu cuidado: o que fica no atendimento e o que o RH vê', publico=f'{COL}; {RH}',
  nec='Vão saber o que eu contei? O que a empresa enxerga?',
  obj='Dar segurança sobre privacidade e explicar os relatórios anonimizados', formato=C, canal=IGLI, collab=FJ,
  msg='O que você conta fica entre você e quem cuida. A empresa recebe só indicadores do grupo, sem identificar ninguém.',
  valid='Confirmar com o time a descrição do fluxo de dados e da LGPD.', status='Pauta proposta',
  criativo='Modelo M3 (cards de interface).'),
 dict(mes=2, sem=8, linha=J, tema='Letramento e workshops', serie='Registro do mês',
  titulo='Por dentro de um workshop nos pilares [tema e data a confirmar]', publico=f'{COL}; {RH}',
  nec='Como é um encontro de letramento na prática?', obj='Registrar uma entrega real do programa',
  formato=V, canal=IGLI, collab=f'{FJ}; empresa-cliente (se autorizar)', msg='A definir com o registro',
  valid='Data, tema e autorização de imagem. Ninguém atendido identificável.', status='Aguardando pauta',
  criativo='Modelo M7.'),
 dict(mes=2, sem=8, linha=P, tema='NR-1 e riscos psicossociais', serie='NR-1 (3/4): cuidado individualizado',
  titulo='Onde o Escritório do Cuidado entra na NR-1', publico=f'{RH}; {SESMT}',
  nec='Isso substitui o SST ou o plano de saúde?',
  obj='Explicar como o Escritório complementa o SST e o plano de saúde na prevenção', formato=C, canal=IGLI, collab=FJ,
  msg='O Escritório do Cuidado complementa o SST e o plano de saúde: leva prevenção e cuidado individualizado para perto de quem trabalha.',
  valid='Responsável médica. Não prometer conformidade com a norma.', status='Pauta proposta', criativo='Modelo M2.'),
 dict(mes=2, sem=8, **time(8, 'Charles', V, 8)),
 dict(mes=2, sem=8, linha=P, tema='Pilar: relações sociais', serie='Pilares da MEV (3: relações sociais)',
  titulo='Cuidar dos vínculos é cuidar da saúde', publico=COL,
  nec='As relações no trabalho afetam a minha saúde?',
  obj='Educar sobre o pilar e conectar com o cuidado em grupo', formato=C, canal=IG, collab=FJ,
  msg='Cuidar dos vínculos é cuidar da saúde.', valid='Responsável médica valida as orientações.',
  status='Pauta proposta', criativo='Modelo M5 (pilar).'),

 # ---------------- MÊS 3 ----------------
 dict(mes=3, sem=9, linha=I, tema='Time', serie='Quem cuida (fechamento)',
  titulo='Oito pessoas, uma pergunta: o que é importante para você?', publico=f'{COL}; {RH}',
  nec='Quem é esse time, junto?', obj='Fechar a série do time e reforçar a pergunta da marca',
  formato=V, canal=IGLI, collab=f'{FJ}; integrantes', msg='Pessoas diferentes, um mesmo jeito de cuidar.',
  valid='Trechos curtos de cada integrante (podem vir dos vídeos da série).', status='Ideia', criativo='Modelo M9.'),
 dict(mes=3, sem=9, linha=P, tema='Pilar: alimentação', serie='Pilares da MEV (4: alimentação)',
  titulo='Comer bem é cuidar do corpo e também da mente', publico=COL,
  nec='Como comer melhor numa rotina corrida, sem dieta impossível?',
  obj='Educar sobre o pilar e mostrar como ele entra no plano', formato=C, canal=IG, collab=FJ,
  msg='Comer bem é cuidar do corpo e também da mente. No plano, a alimentação parte da sua rotina real.',
  valid='Responsável médica. Nenhum número da cartilha sem fonte verificada.', status='Ideia', criativo='Modelo M5.'),
 dict(mes=3, sem=9, linha=J, tema='Depoimentos', serie='Da intenção à evidência (3: adesão entre colegas)',
  titulo='"Ela não gosta de gravar vídeo. Gravou um mesmo assim."', publico=f'{COL}; {RH}',
  nec='Os colegas recomendam?', obj='Mostrar adesão espontânea com trechos aprovados', formato=V, canal=IGLI, collab=FJ,
  msg='"O que é bom a gente tem que compartilhar."',
  valid='Trechos A1 a A6 aprovados; o vídeo original não é postado. Formato confirmado com a responsável médica (CFM).',
  status='Ideia', criativo='Modelo M8 animado (trechos em cards).'),
 dict(mes=3, sem=9, linha=I, tema='Quem cuida de quem cuida', serie='',
  titulo='Quem cuida também precisa ser cuidado', publico=f'{RH}; {SESMT}',
  nec='Por que a Fair Health fala também de quem cuida?',
  obj='Mostrar a visão integrada do cuidado (dos Aims da saúde à empresa)', formato=C, canal=IGLI, collab=FJ,
  msg='Cuidado bom é bom para quem recebe e para quem cuida.',
  valid='Conceitos e datas dos Aims (IHI) com fonte.', status='Ideia', criativo='Modelo M2.'),
 dict(mes=3, sem=10, linha=P, tema='NR-1 e riscos psicossociais', serie='NR-1 (4/4): acompanhamento',
  titulo='A NR-1 não termina no diagnóstico: como acompanhar ao longo do ano', publico=f'{RH}; {SESMT}',
  nec='Fiz o diagnóstico. Como sei se o que estamos fazendo funciona?',
  obj='Mostrar acompanhamento contínuo e indicadores', formato=V, canal=IGLI, collab=FJ,
  msg='Cuidado sem dado é intenção. O acompanhamento mostra o que muda ao longo do programa.',
  valid='Responsável médica. Posição [SUGESTÃO] 4 do voice.md precisa de aprovação.', status='Ideia', criativo='Modelo M9.'),
 dict(mes=3, sem=10, linha=J, tema='Case e evento', serie='Da intenção à evidência (4: experiência)',
  titulo='Onde o modelo começou: de 34 para cerca de 20 reclamações por mês', publico=f'{RH}; {SESMT}',
  nec='Esse modelo já funcionou em algum lugar?',
  obj='Mostrar o case do Hospital Dia Santo Amaro, apresentado em congresso', formato=C, canal=IGLI, collab=FJ,
  msg='No Hospital Dia Santo Amaro, as reclamações mensais caíram de 34 para cerca de 20, com 98,2% de conformidade média dos processos.',
  valid='Dados liberados (voice.md). Citar o congresso com nome e ano corretos.', status='Ideia', criativo='Modelo M6.'),
 dict(mes=3, sem=10, linha=I, tema='Relacionamento', serie='',
  titulo='O que é importante para você em 2027?', publico=f'{COL}; {RH}',
  nec='(Convite à participação) O que eu quero cuidar no próximo ano?',
  obj='Gerar conversa e coletar respostas para pautas futuras', formato=V, canal=IG, collab=FJ,
  msg='Todo cuidado começa com uma pergunta. Esta é a nossa para o seu próximo ano.',
  valid='Caixa de perguntas nos Stories antes da gravação.', status='Ideia', criativo='Modelo M9 com respostas do time.'),
 dict(mes=3, sem=10, linha=P, tema='Indicadores', serie='Da intenção à evidência (o conjunto)',
  titulo='Indicadores que o RH deve acompanhar: adesão, retorno, absenteísmo, presenteísmo, clima e sinistralidade', publico=f'{RH}; {CFO}',
  nec='Quais números mostram que o cuidado está funcionando?',
  obj='Educar o decisor sobre indicadores e posicionar o relatório', formato=C, canal=IGLI, collab=FJ,
  msg='Adesão é o primeiro resultado a medir. Os outros indicadores vêm ao longo do programa.',
  valid='Definição de cada indicador. Posição [SUGESTÃO] 3 do voice.md precisa de aprovação.', status='Ideia', criativo='Modelo M6.'),
 dict(mes=3, sem=11, linha=J, tema='Entregas do ano', serie='Registro do mês',
  titulo='2026 em entregas [dados a confirmar]', publico=f'{RH}; {COL}',
  nec='O que a Fair Health já entregou?', obj='Mostrar movimento e prova social', formato=V, canal=IGLI,
  collab=f'{FJ}; clientes (se autorizarem)', msg='A definir com os dados liberados do ano',
  valid='Usar somente números liberados.', status='Aguardando pauta', criativo='Modelo M7.'),
 dict(mes=3, sem=11, linha=I, tema='Por que existimos', serie='Da intenção à evidência (afastamento)',
  titulo='O afastamento não começa no atestado', publico=f'{RH}; {CFO}',
  nec='Por que o RH só percebe o problema quando já virou afastamento?',
  obj='Explicar por que a Fair Health existe e apresentar o argumento da prevenção', formato=C, canal=IGLI, collab=FJ,
  msg='O adoecimento começa antes do atestado, na rotina. A Fair Health existe para cuidar antes.',
  valid='Arte existente (instagram/posts/estreia, #1) a adaptar ao visual atual. HBR 2022 (−28% de absenteísmo) só com a fonte no slide.',
  status='Ideia', criativo='Modelo M6 no slide do dado.'),
 dict(mes=3, sem=11, linha=P, tema='Sinistralidade', serie='',
  titulo='Sinistralidade alta não se resolve trocando de plano', publico=f'{RH}; {CFO}',
  nec='O reajuste do plano veio alto de novo. O que dá para fazer além de trocar de operadora?',
  obj='Posicionar a Fair Health na causa raiz, com opinião', formato=V, canal=IGLI, collab=FJ,
  msg='Sinistralidade se resolve tratando a causa antes que ela vire conta.',
  valid='Posição [SUGESTÃO] 2 do voice.md precisa de aprovação. Não prometer redução de custo; SHRM 2023 (−22%) só com a fonte.',
  status='Ideia', criativo='Modelo M9 (opinião para a câmera).'),
 dict(mes=3, sem=11, linha=J, tema='Eventos', serie='',
  titulo='Do palco do TEDx: felicidade e saúde mental no trabalho', publico=RH,
  nec='Quem pensa a Fair Health e o que ela defende?',
  obj='Trazer a autoridade do fundador a partir de um evento público', formato=C, canal=IGLI,
  collab=f'{FJ}; Fernando Brancaccio', msg='A definir a partir da palestra',
  valid='Confirmar conteúdo, link e direito de uso de trechos da palestra.', status='Ideia', criativo='Modelo M8 (citações da palestra).'),
 dict(mes=3, sem=12, linha=I, tema='Crenças', serie='',
  titulo='Fazemos o certo pelos motivos certos: o que isso quer dizer na prática', publico=f'{RH}; {COL}',
  nec='No que a Fair Health acredita?', obj='Reforçar valores com exemplos reais do jeito de cuidar',
  formato=V, canal=IGLI, collab=FJ, msg='O retorno é consequência, não ponto de partida.',
  valid='Exemplos reais escolhidos com o time.', status='Ideia', criativo='Modelo M9.'),
 dict(mes=3, sem=12, linha=P, tema='Pilar: propósito', serie='Pilares da MEV (5: propósito e espiritualidade)',
  titulo='Propósito também é pilar de saúde', publico=COL,
  nec='O que propósito tem a ver com saúde?', obj='Educar sobre o pilar e conectar com a pergunta da marca',
  formato=C, canal=IG, collab=FJ, msg='Saber o que é importante para você também é cuidar da saúde.',
  valid='Responsável médica valida a abordagem do pilar.', status='Ideia', criativo='Modelo M5.'),
 dict(mes=3, sem=12, linha=J, tema='Registro do mês', serie='Registro do mês',
  titulo='[Registro do mês: workshop, evento ou entrega, a confirmar]', publico='A definir',
  nec='O que a Fair Health está fazendo agora?', obj='Mostrar movimento real e prova social',
  formato=V, canal=IGLI, collab=f'{FJ}; parceiros envolvidos (se autorizarem)', msg='A definir',
  valid='Agenda de eventos e entregas do mês. Autorizações de imagem.', status='Aguardando pauta', criativo='Modelo M7.'),
 dict(mes=3, sem=12, linha=P, tema='Oferta', serie='',
  titulo='Diagnóstico gratuito: o primeiro passo para levar o Escritório do Cuidado à sua empresa', publico=f'{RH}; {CFO}',
  nec='Como eu levo isso para a minha empresa?', obj='Converter com uma chamada comercial clara',
  formato=C, canal=IGLI, collab=FJ, msg='Comente CUIDADO e enviamos como funciona o diagnóstico gratuito.',
  valid='Confirmar o destino do link (site) e quem responde os comentários e DMs.', status='Ideia',
  criativo='Modelo M1 (capa) e M2 (passos).'),
]

# checagens: 16 por mês, 4 por semana, sem linha repetida em sequência, vídeo intercalado
assert len(ROWS) == 48
for k in range(1, len(ROWS)):
    assert ROWS[k]['linha'] != ROWS[k - 1]['linha'], (k, ROWS[k]['titulo'])
for k, r in enumerate(ROWS):
    eh_video = r['formato'] == V
    assert eh_video == (k % 2 == 0), (k, r['titulo'], r['formato'])
    assert r['sem'] == k // 4 + 1, (k, r['titulo'])

COLS = [
    ('ID', 8), ('Mês', 8), ('Semana', 8), ('Data sugerida', 12), ('Linha editorial', 16), ('Tema', 20),
    ('Série (parte)', 24), ('Título provisório', 36), ('Público', 22), ('Necessidade ou pergunta do público', 34),
    ('Objetivo', 30), ('Formato', 11), ('Canal', 18), ('Collab / perfis', 24), ('Mensagem principal', 44),
    ('Validação necessária', 40), ('Status', 16), ('Texto / legenda', 50), ('Roteiro / estrutura', 50),
    ('Criativo (modelo e arte)', 40), ('Responsável', 14), ('Comentários da revisão', 30),
]
KEYS = [None, None, None, None, 'linha', 'tema', 'serie', 'titulo', 'publico', 'nec', 'obj', 'formato',
        'canal', 'collab', 'msg', 'valid', 'status', 'texto', 'roteiro', 'criativo', None, None]

wb = Workbook()

# ---------- Como usar ----------
ws0 = wb.active
ws0.title = 'Como usar'
linhas = [
    ('Calendário editorial · Fair Health · Instagram e LinkedIn', TITULO),
    ('Versão de 08/10/2026. Públicos e territórios são provisórios até chegar o material estratégico complementar.', NORMAL),
    ('', NORMAL),
    ('Como trabalhamos nesta planilha', CAB),
    ('1. Revisão mensal (conjunta): antes de cada mês começar, revisamos juntos temas, títulos, públicos, linhas e datas do mês seguinte. Pauta aprovada muda de "Pauta proposta" para "Pauta aprovada".', NORMAL),
    ('2. Detalhamento semanal: a cada semana, detalhamos texto, roteiro e criativo das peças da semana seguinte (colunas R a T) e passamos pela validação.', NORMAL),
    ('3. Comentários: use a coluna "Comentários da revisão" (V) ou comentários de célula. Não apague pautas descartadas: mude o status para "Descartada".', NORMAL),
    ('4. Datas: a data de início (célula amarela na aba Calendário) recalcula todas as datas sugeridas. Para mudar uma data só, digite a nova data por cima.', NORMAL),
    ('', NORMAL),
    ('Cadência proposta (a validar)', CAB),
    ('4 publicações por semana no feed, alternando vídeo e estático/carrossel (2 vídeos por semana). Os vídeos circulam entre as três linhas editoriais.', NORMAL),
    ('As três linhas entram em rodízio: duas peças seguidas nunca são da mesma linha. A aba Visão mensal conta o equilíbrio de cada mês.', NORMAL),
    ('Os dias da semana NÃO estão definidos: as datas são só uma distribuição sugerida (evitando feriados nacionais) e devem ser ajustadas na revisão mensal.', NORMAL),
    ('Stories ficam fora da contagem e servem de apoio: repost, caixa de perguntas e bastidores.', NORMAL),
    ('', NORMAL),
    ('Marcações', CAB),
    ('[PENDENTE] = informação que ainda precisa ser confirmada. [SUGESTÃO] = posição ou texto proposto que precisa de aprovação.', NORMAL),
    ('Coluna "Validação necessária": tudo o que precisa de aval antes de publicar. Afirmações clínicas, benefícios, prazos e resultados passam pela responsável médica.', NORMAL),
    ('Só usamos números liberados: Hospital Dia Santo Amaro, Chemitec (1º semestre de 2026), HBR 2022 e SHRM 2023 com a fonte (instagram/voice.md).', NORMAL),
    ('', NORMAL),
    ('Fluxo de status', CAB),
    ('Ideia → Pauta proposta → Pauta aprovada → Aguardando material → Roteiro pronto → Em produção → Em revisão → Validação médica → Aprovado → Agendado → Publicado (ou Descartada / Aguardando pauta)', NORMAL),
    ('', NORMAL),
    ('Abas', CAB),
    ('Calendário: todas as pautas, uma por linha. O Mês 1 está completo no nível mensal e a semana 1 já tem texto, roteiro e criativo. Meses 2 e 3 estão no nível de tema e título.', NORMAL),
    ('Visão mensal: foco, temas por linha, séries, datas do calendário e o equilíbrio de cada mês (contagens automáticas).', NORMAL),
    ('Linhas e públicos: o que entra em cada linha, os públicos provisórios com as perguntas de cada um e os territórios temáticos.', NORMAL),
    ('Séries: assuntos amplos desdobrados em peças curtas e complementares.', NORMAL),
    ('Pendências: o que precisa ser confirmado e quais pautas cada item trava.', NORMAL),
    ('Listas: valores das listas suspensas.', NORMAL),
    ('', NORMAL),
    ('Documentos de apoio no repositório: instagram/estrategia.md (estratégia), instagram/voice.md (voz), instagram/diretrizes-criativas.md (visual e formatos).', NORMAL),
]
for i, (t, f) in enumerate(linhas, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws0.column_dimensions['A'].width = 140

# ---------- Calendário ----------
ws = wb.create_sheet('Calendário')
ws['A1'] = 'Calendário editorial · Fair Health'
ws['A1'].font = TITULO
ws['A2'] = 'Início do Mês 1 (segunda-feira):'
ws['A2'].font = CAB
ws.merge_cells('A2:C2')
ws['D2'] = dt.date(2026, 10, 19)
ws['D2'].number_format = 'DD/MM/YYYY'
ws['D2'].fill = FILL_INPUT
ws['D2'].font = Font(name=F, bold=True, color='0000FF', size=10)
ws['D2'].comment = Comment('Data proposta para começar o Mês 1 (sugestão da revisão de 08/10). Mude aqui e todas as datas sugeridas se ajustam.', 'Planejamento')
ws['E2'] = 'Mude a data amarela para reposicionar todo o calendário. Os dias da semana são só sugestão.'
ws['E2'].font = Font(name=F, italic=True, size=9)
HR = 4
for j, (nome, larg) in enumerate(COLS, 1):
    c = ws.cell(row=HR, column=j, value=nome)
    c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(j)].width = larg
ws.row_dimensions[HR].height = 30

cont = {}
for k, r in enumerate(ROWS):
    row = HR + 1 + k
    cont[r['mes']] = cont.get(r['mes'], 0) + 1
    pos = k % 4
    ws.cell(row=row, column=1, value=f"M{r['mes']}-{cont[r['mes']]:02d}")
    ws.cell(row=row, column=2, value=f"Mês {r['mes']}")
    ws.cell(row=row, column=3, value=r['sem'])
    d = ws.cell(row=row, column=4, value=f'=$D$2+(C{row}-1)*7+{OFF[r["sem"]][pos]}')
    d.number_format = 'DD/MM/YYYY'
    for j, key in enumerate(KEYS, 1):
        if key:
            ws.cell(row=row, column=j, value=r.get(key, '') or None)
    for j in range(1, len(COLS) + 1):
        c = ws.cell(row=row, column=j)
        c.font = NORMAL if j != 4 else Font(name=F, size=10)
        c.alignment = WRAP
        c.border = BORDA
last = HR + len(ROWS)
ws.freeze_panes = 'I5'
ws.auto_filter.ref = f'A{HR}:{get_column_letter(len(COLS))}{last + 150}'

# formatação condicional por linha editorial (coluna E)
for valor, cor in ((I, 'D7EAE5'), (P, 'F7EBC5'), (J, 'DFE4F0')):
    ws.conditional_formatting.add(f'E{HR + 1}:E{last + 150}',
        FormulaRule(formula=[f'$E{HR + 1}="{valor}"'], fill=PatternFill('solid', fgColor=cor)))
ws.conditional_formatting.add(f'L{HR + 1}:L{last + 150}',
    FormulaRule(formula=[f'$L{HR + 1}="{V}"'], font=Font(name=F, bold=True, color='1F3A36')))

# ---------- Listas ----------
wl = wb.create_sheet('Listas')
LISTAS = {
    'A': ('Linha editorial', [I, P, J]),
    'B': ('Formato', [V, C, FU, 'Stories']),
    'C': ('Canal', [IGLI, IG, 'Instagram Stories', 'LinkedIn']),
    'D': ('Status', ['Ideia', 'Pauta proposta', 'Pauta aprovada', 'Aguardando material', 'Aguardando pauta',
                     'Roteiro pronto', 'Em produção', 'Em revisão', 'Validação médica', 'Aprovado', 'Agendado',
                     'Publicado', 'Descartada']),
    'E': ('Mês', [f'Mês {n}' for n in range(1, 13)]),
    'F': ('Públicos (provisórios)', [RH, COL, SESMT, CFO]),
}
for col, (tit, vals) in LISTAS.items():
    wl[f'{col}1'] = tit
    wl[f'{col}1'].font, wl[f'{col}1'].fill = CAB, FILL_CAB
    for i, v in enumerate(vals, 2):
        wl[f'{col}{i}'] = v
        wl[f'{col}{i}'].font = NORMAL
    wl.column_dimensions[col].width = 30

def dv(col_lista, n, alvo):
    v = DataValidation(type='list', formula1=f"=Listas!${col_lista}$2:${col_lista}${n + 1}", allow_blank=True)
    v.error = 'Escolha um valor da lista (aba Listas). Para criar um valor novo, inclua-o na aba Listas.'
    v.errorStyle = 'warning'
    ws.add_data_validation(v)
    v.add(alvo)
dv('A', 3, f'E{HR + 1}:E{last + 150}')
dv('B', 4, f'L{HR + 1}:L{last + 150}')
dv('C', 4, f'M{HR + 1}:M{last + 150}')
dv('D', 13, f'Q{HR + 1}:Q{last + 150}')
dv('E', 12, f'B{HR + 1}:B{last + 150}')

# ---------- Visão mensal ----------
wv = wb.create_sheet('Visão mensal', 1)
wv['A1'] = 'Visão mensal: temas e títulos'
wv['A1'].font = TITULO
wv['A2'] = 'Primeiro nível do planejamento. Os títulos de cada mês estão na aba Calendário. As contagens são automáticas e servem para a checagem de equilíbrio na revisão mensal.'
wv['A2'].font = Font(name=F, italic=True, size=9)
cab = ['Mês', 'Período', 'Foco do mês', 'Institucional', 'Produto', 'Projetos, eventos e entregas',
       'Séries em andamento', 'Datas do calendário a avaliar', 'Total', 'Institucional (nº)', 'Produto (nº)',
       'Projetos (nº)', 'Vídeos', 'Estáticos e carrosséis', 'Pautas com validação médica']
larg = [8, 16, 34, 34, 40, 36, 34, 36, 8, 12, 10, 11, 9, 13, 14]
for j, (t, w) in enumerate(zip(cab, larg), 1):
    c = wv.cell(row=4, column=j, value=t)
    c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
    c.alignment = Alignment(wrap_text=True, vertical='center')
    wv.column_dimensions[get_column_letter(j)].width = w
MESES = [
    ('Mês 1', 1, 'Apresentar quem somos e como o cuidado funciona, com acolhimento (sem argumento de retorno).',
     'Boas-vindas; time (Marília, Nina, Ronald, Floriana); origem: do Escritório do Paciente ao Escritório do Cuidado.',
     'Como funciona o Escritório do Cuidado; MEV e os 7 pilares; Plano Individualizado de Cuidado; pilar sono; wearables.',
     'Um dia no Escritório; workshops nos pilares; relatório para o RH; como o Escritório chega a uma empresa; registro do mês.',
     'Como funciona o cuidado (1 a 3); Quem cuida (1 a 4); Pilares da MEV (abertura e sono); Como chega a uma empresa (1).',
     '02/11 feriado (Finados). 14/11 Dia Mundial do Diabetes: avaliar com a responsável médica um conteúdo de alimentação e atividade física, sem promessa clínica.'),
    ('Mês 2', 5, 'Aprofundar: NR-1 em série, saúde mental, privacidade dos dados e a primeira prova (Chemitec).',
     'Time (Renata, Liliane, Fernando, Charles); ecossistema FairJob.',
     'NR-1 (1 a 3) e versão para colaboradores; pilares saúde mental e relações sociais; seus dados, seu cuidado.',
     'Bastidores de uma implantação; Chemitec 1º semestre de 2026; depoimento B1; workshop nos pilares (registro do mês).',
     'NR-1 (1 a 3); Quem cuida (5 a 8); Pilares da MEV (2 e 3); Da intenção à evidência (1 e 2); Como chega a uma empresa (2).',
     'Novembro Azul: avaliar se há conexão real antes de usar. 20/11 feriado (Consciência Negra).'),
    ('Mês 3', 9, 'Evidência e convite: indicadores um a um, case HDSA, por que existimos e o diagnóstico gratuito.',
     'Time junto; quem cuida de quem cuida; o que é importante para você em 2027; por que existimos; fazemos o certo pelos motivos certos.',
     'NR-1 (4); indicadores para o RH; sinistralidade (opinião); pilares alimentação e propósito; diagnóstico gratuito.',
     'Depoimento em vídeo de cards; HDSA; 2026 em entregas; TEDx; registro do mês.',
     'NR-1 (4); Pilares da MEV (4 e 5); Da intenção à evidência (3, 4 e conjunto).',
     'Festas de fim de ano (25/12 e 01/01 feriados): reduzir a cadência se a produção apertar.'),
]
for i, (m, sem0, foco, ti, tp, tj, series, datas) in enumerate(MESES, 5):
    ini = "'Calendário'!$D$2"
    vals = [m, f'=TEXT({ini}+{(sem0 - 1) * 7},"DD/MM")&" a "&TEXT({ini}+{(sem0 - 1) * 7 + 27},"DD/MM")',
            foco, ti, tp, tj, series, datas]
    for j, v in enumerate(vals, 1):
        wv.cell(row=i, column=j, value=v)
    R = "'Calendário'"
    wv.cell(row=i, column=9, value=f'=COUNTIFS({R}!$B$5:$B$400,$A{i})')
    wv.cell(row=i, column=10, value=f'=COUNTIFS({R}!$B$5:$B$400,$A{i},{R}!$E$5:$E$400,"{I}")')
    wv.cell(row=i, column=11, value=f'=COUNTIFS({R}!$B$5:$B$400,$A{i},{R}!$E$5:$E$400,"{P}")')
    wv.cell(row=i, column=12, value=f'=COUNTIFS({R}!$B$5:$B$400,$A{i},{R}!$E$5:$E$400,"{J}")')
    wv.cell(row=i, column=13, value=f'=COUNTIFS({R}!$B$5:$B$400,$A{i},{R}!$L$5:$L$400,"{V}")')
    wv.cell(row=i, column=14, value=f'=I{i}-M{i}')
    wv.cell(row=i, column=15, value=f'=SUMPRODUCT(({R}!$B$5:$B$400=$A{i})*ISNUMBER(SEARCH("responsável médica",{R}!$P$5:$P$400)))')
    for j in range(1, 16):
        c = wv.cell(row=i, column=j)
        c.font, c.alignment, c.border = NORMAL, WRAP, BORDA
wv.freeze_panes = 'C5'

# ---------- Linhas e públicos ----------
wp = wb.create_sheet('Linhas e públicos', 2)
def tabela(ws_, r0, titulo, cab_, dados, larguras=None):
    ws_.cell(row=r0, column=1, value=titulo).font = Font(name=F, bold=True, size=12, color=TEXTO)
    for j, t in enumerate(cab_, 1):
        c = ws_.cell(row=r0 + 1, column=j, value=t)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, linha in enumerate(dados, r0 + 2):
        for j, v in enumerate(linha, 1):
            c = ws_.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = NORMAL, WRAP, BORDA
    return r0 + 2 + len(dados) + 1

for j, w in enumerate([24, 46, 46, 40, 30], 1):
    wp.column_dimensions[get_column_letter(j)].width = w
r = tabela(wp, 1, 'Linhas editoriais', ['Linha', 'O que entra', 'Objetivos típicos', 'Formatos que funcionam', 'Peso no mês'], [
    (I, 'O que é a Fair Health, por que existe, de onde veio (Escritório do Paciente → Escritório do Cuidado, ecossistema FairJob), em que acredita e quem compõe o time, com a apresentação individual de cada integrante.',
     'Apresentar, humanizar, gerar confiança.', 'Vídeo com o time, retratos, carrossel de origem.', 'Cerca de 1/3'),
    (P, 'Escritório do Cuidado (espaço, percurso, sigilo), Medicina do Estilo de Vida e os 7 pilares, Plano Individualizado de Cuidado e tecnologia (smartband, bioimpedância, plataforma, relatórios anonimizados). Temas do entorno (NR-1, bem-estar, saúde mental) entram aqui quando a ligação com a oferta é clara.',
     'Explicar, educar com rigor, tirar dúvidas, levar à conversa comercial.', 'Carrossel explicativo, vídeo de profissional para a câmera.', 'Cerca de 1/3 a 2/5'),
    (J, 'O que a Fair Health está fazendo e entregando: operação nas empresas (bastidores), workshops e letramento, cases e números liberados, eventos e participações, entregas para o RH.',
     'Provar, mostrar movimento, gerar prova social.', 'Vídeo de bastidor, carrossel de case, registro de evento.', 'Cerca de 1/4 a 1/3'),
])
r = tabela(wp, r, 'Públicos (PROVISÓRIO: confirmar com o material estratégico complementar)',
           ['Público', 'Quem é (como está documentado)', 'Perguntas e necessidades', 'Como falar', 'Cuidados'], [
    (RH, 'Público principal: CHROs, gerentes de pessoas e business partners de empresas de médio e grande porte, principalmente indústria.',
     'Como atender à NR-1 de um jeito que mude algo? Meu programa de saúde tem adesão? O que eu recebo e o que não posso ver? Como começar?',
     'Direto, com processo e evidência, sem jargão clínico.', 'Números só liberados e com fonte.'),
    (COL, 'Quem é ou pode ser atendido pelo Escritório do Cuidado (público dos primeiros posts de acolhimento).',
     'Como funciona? Vou ser julgado? O que eu contar fica entre nós? Cabe na minha rotina?',
     'Acolhedor, em segunda pessoa, prático, sem números de negócio.', 'Nunca expor quem é atendido.'),
    (SESMT, 'Público secundário: médicos do trabalho e equipes de SST.',
     'Isso substitui ou complementa o que já fazemos? Como conversa com o SST e com a NR-1?',
     'Técnico, com fontes e precisão.', 'Validação da responsável médica em tudo o que for clínico ou normativo.'),
    (CFO, 'Público secundário: diretoria financeira.',
     'Qual o retorno? Como medir? O que muda na sinistralidade?',
     'Indicadores com fonte, sem promessa.', 'Argumento de retorno só a partir do Mês 2, um indicador por vez.'),
])
r = tabela(wp, r, 'Territórios temáticos (PROVISÓRIO)',
           ['Território', 'Ligação real com a Fair Health', 'Como adaptar ao público', 'Validação', 'Onde aparece'], [
    ('NR-1 e riscos psicossociais', 'O Executive Summary posiciona a oferta como aderente à NR-1 (programa de prevenção de riscos psicossociais e plataforma de indicadores).',
     'RH: o que muda na gestão. SESMT: como complementa o SST. Colaboradores: seu bem-estar no trabalho também é assunto da empresa.',
     'Texto da norma com fonte oficial; responsável médica; não prometer conformidade.', 'Série NR-1 (Meses 2 e 3)'),
    ('Bem-estar e saúde no trabalho', 'Cuidado preventivo dentro da empresa; adesão como primeiro resultado.',
     'RH: adesão e experiência. Colaboradores: cuidado perto, sem fila.', 'Não chamar de "app de bem-estar" nem de "benefício" como mimo.', 'Meses 1 a 3'),
    ('Os 7 pilares da MEV', 'Base do Plano Individualizado de Cuidado (cartilha do Escritório do Cuidado).',
     'Colaboradores: hábito possível. RH: o que o ambiente da empresa pode facilitar.', 'Responsável médica; números da cartilha sem fonte ficam de fora.', 'Série Pilares, um pilar por vez ao longo dos meses'),
    ('Saúde mental', 'Psicólogos no Escritório; guia "Saúde mental: o cuidado no dia a dia" (OMS, APA, IPq-USP).',
     'Colaboradores: acolher e dizer onde pedir ajuda. RH: liderança e clima.', 'Sempre fechar com CVV 188; responsável médica.', 'Meses 2 e 3'),
    ('Tecnologia e dados', 'Smartband, bioimpedância, plataforma, relatórios anonimizados integrados ao fairdata.',
     'Colaboradores: privacidade. RH: o que recebe e o que não vê.', 'Nunca dado individual; confirmar fluxo de dados.', 'Série Tecnologia a favor da pessoa'),
    ('Experiência do colaborador e cuidado centrado na pessoa', 'Origem no Escritório do Paciente e nos Aims da saúde (IHI).',
     'RH e SESMT: modelo com história e método.', 'Datas e conceitos com fonte; [PENDENTE] quem idealizou.', 'Origem (Mês 1), Quem cuida de quem cuida (Mês 3)'),
    ('Indicadores e retorno', 'Provas liberadas: HDSA, Chemitec, HBR 2022, SHRM 2023.',
     'RH e diretoria financeira: um indicador por vez.', 'Só números liberados, sempre com fonte.', 'Série Da intenção à evidência (Meses 2 e 3)'),
    ('Datas e temas do entorno', 'Só quando a data conversa com um pilar ou com o serviço.',
     'Escolher o público que a data realmente envolve.', 'Responsável médica; evitar oportunismo.', 'Coluna "Datas do calendário a avaliar" da Visão mensal'),
])

# ---------- Séries ----------
wsr = wb.create_sheet('Séries', 3)
for j, w in enumerate([30, 50, 60, 16, 44], 1):
    wsr.column_dimensions[get_column_letter(j)].width = w
tabela(wsr, 1, 'Séries: assuntos amplos em peças curtas e complementares',
       ['Série', 'Lógica da sequência', 'Partes (ID no Calendário)', 'Meses', 'Observações'], [
    ('Como funciona o cuidado', 'Contexto e percurso → cuidado individualizado → acompanhamento.',
     'M1-02 O percurso · M1-08 O Plano Individualizado de Cuidado · M1-14 O acompanhamento (wearables)', 'Mês 1',
     'Detalhes técnicos validados pela responsável médica.'),
    ('NR-1', 'Contexto → diagnóstico → cuidado individualizado → acompanhamento, com uma versão para colaboradores.',
     'M2-02 Contexto · M2-07 Versão para colaboradores · M2-09 Diagnóstico · M2-14 Cuidado individualizado · M3-05 Acompanhamento', 'Meses 2 e 3',
     'Fonte oficial para o texto da norma. Não prometer conformidade.'),
    ('Pilares da MEV', 'Um pilar por vez, espalhado pelo calendário. Cada pilar pode ter até 3 peças em momentos diferentes: por que importa (colaboradores) → como entra no plano (produto) → o que a empresa pode facilitar (RH).',
     'M1-05 Abertura (7 pilares) · M1-11 Sono · M2-04 Saúde mental e estresse · M2-16 Relações sociais · M3-02 Alimentação · M3-14 Propósito · Atividade física e evitar substâncias no Mês 4',
     'Meses 1 a 4', 'Um post não precisa esgotar o pilar. Retomar os pilares com outros ângulos depois.'),
    ('Quem cuida', 'Uma pessoa do time por semana, sem cargo, com a resposta a "O que é importante para você?". Fecha com o time junto.',
     'M1-04 Marília · M1-07 Nina · M1-13 Ronald · M1-16 Floriana · M2-03 Renata · M2-05 Liliane · M2-11 Fernando · M2-15 Charles · M3-01 Time junto', 'Meses 1 a 3',
     'Ordem a confirmar com o time. Formato (vídeo ou retrato) conforme a preferência de cada pessoa.'),
    ('Tecnologia a favor da pessoa', 'O que a tecnologia mede → o que fica no atendimento → o que o RH recebe.',
     'M1-14 Wearables · M1-09 Relatório para o RH · M2-12 Seus dados, seu cuidado', 'Meses 1 e 2', 'Nunca dado individual.'),
    ('Da intenção à evidência', 'Um indicador por vez, depois o conjunto (decisão de 01/10: o argumento de retorno começa depois da estreia).',
     'M2-06 Adesão e retorno (Chemitec) · M2-10 Experiência (depoimento B1) · M3-03 Adesão entre colegas · M3-06 HDSA · M3-08 Indicadores · M3-10 Afastamento', 'Meses 2 e 3',
     'Só números liberados, sempre com fonte.'),
    ('Como chega a uma empresa', 'Explicação das etapas → bastidor de uma implantação real.',
     'M1-12 Diagnóstico, proposta e implantação · M2-01 Bastidores de uma implantação', 'Meses 1 e 2',
     'Deixa clara a oferta comercial. Não citar prazos sem confirmação.'),
    ('Registro do mês', 'Uma peça por mês sobre um projeto, evento ou entrega real.',
     'M1-15 · M2-13 (workshop) · M3-09 (2026 em entregas) · M3-15', 'Todos os meses', 'Depende da agenda de eventos e entregas, que ainda não está documentada.'),
])

# ---------- Pendências ----------
wpd = wb.create_sheet('Pendências', 4)
for j, w in enumerate([5, 60, 30, 40, 14], 1):
    wpd.column_dimensions[get_column_letter(j)].width = w
pend = [
    ('Material estratégico complementar: confirmar públicos e aprofundar os territórios temáticos.', 'Time Fair Health', 'Colunas Público e Necessidade de todas as pautas'),
    ('Nome da responsável médica que valida o conteúdo clínico (o material anterior citava o Dr. Ronald).', 'Time Fair Health', 'Todas as pautas com "responsável médica" na validação'),
    ('Assistentes sociais fazem parte do atendimento? (o Executive Summary cita médicos e psicólogos)', 'Time Fair Health', 'M1-01 e textos que citam o time'),
    ('Etapas e prazos clínicos: bioimpedância, smartband por 48 h, retorno, programa de 12 meses.', 'Responsável médica', 'M1-02, M1-08, M1-14'),
    ('Nomes dos 7 pilares: usar a lista da cartilha? (a apresentação institucional traz 6)', 'Responsável médica', 'Série Pilares da MEV'),
    ('Texto, prazos e obrigações da NR-1 com fonte oficial.', 'Responsável médica + fonte oficial', 'Série NR-1'),
    ('Collab no Instagram: perfil da FairJob, quem envia e quem aceita o convite, se todos os posts vão em collab, perfis dos integrantes, limite de colaboradores por post.', 'Comunicação + FairJob', 'Todas as pautas (coluna Collab)'),
    ('LinkedIn pela página da FairJob: quem publica, frequência, se a Fair Health terá página própria, se os integrantes compartilham, adaptação das legendas.', 'Comunicação + FairJob', 'Todas as pautas com canal LinkedIn'),
    ('Fotos e vídeos reais do espaço (as imagens atuais parecem renders) e autorização da empresa para gravar.', 'Time Fair Health', 'M1-01, M1-03, M2-01'),
    ('Material de cada integrante: vídeo curto ou retrato, descrição sem cargo e resposta à pergunta. Ordem da série.', 'Time Fair Health', 'Série Quem cuida'),
    ('Workshops e letramento: formato, frequência, temas e próximas datas.', 'Time Fair Health', 'M1-06, M2-13'),
    ('Agenda de eventos, projetos e entregas para o Registro do mês.', 'Time Fair Health', 'M1-15, M2-13, M3-09, M3-15'),
    ('Cores secundárias e elementos visuais: todos em teste. O amarelo usado no carrossel de acolhimento conflita com uma nota anterior (amarelo é cor da Fair Edu).', 'Comunicação + sócios', 'Todas as artes'),
    ('Raleway é a fonte oficial da marca? Logo em alta resolução e PNG do coração com fundo transparente.', 'Comunicação', 'Todas as artes'),
    ('Destino do link e da chamada CUIDADO (www.fairhealth.com.br) e quem responde comentários e DMs.', 'Comunicação', 'M1-12, M3-16'),
    ('Limites de publicidade médica (CFM) e formato dos depoimentos.', 'Responsável médica', 'M2-10, M3-03'),
    ('Chemitec: autorização para citar ou marcar a empresa no post.', 'Time Fair Health', 'M2-06'),
    ('Posições [SUGESTÃO] 1 a 4 do voice.md (NR-1, sinistralidade, adesão, dado).', 'Sócios', 'M2-02, M3-05, M3-08, M3-11'),
    ('Quem idealizou o Escritório do Paciente.', 'Time Fair Health', 'M1-10'),
]
tabela(wpd, 1, 'Pendências e validações', ['#', 'Pendência', 'Quem confirma', 'Pautas afetadas', 'Situação'],
       [(i, a, b, c, 'Aberta') for i, (a, b, c) in enumerate(pend, 1)])

wb.save(OUT)
print('ok', OUT)
