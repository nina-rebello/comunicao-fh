# Versão simples (pedido da Nina, 14/10): uma aba, uma tabela, uma linha por post, só as colunas principais.
# Conteúdo em conteudo_posts_19-10_a_13-11.py. A versão completa (8 abas) é gerada por gerar_cronograma_posts.py.
# Uso: python3 marketing/ferramentas/gerar_cronograma_simples.py marketing/Cronograma_Posts_19-10_a_13-11_simples.xlsx
import math, os, runpy, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

C = runpy.run_path(os.path.join(os.path.dirname(__file__), 'conteudo_posts_19-10_a_13-11.py'))
DIAS = ['segunda', 'terça', 'quarta', 'quinta', 'sexta', 'sábado', 'domingo']
COR_FMT = {'Reel': '96C5BD', 'Carrossel': 'BFE0D9', 'Foto única': 'F0F0E9'}
fino = Side(style='thin', color='C9D6D2')
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)

def roteiro(p):
    linhas = []
    for parte, texto, fala, visual, obs in p['roteiro']:
        t = texto.replace('\n', ' / ')
        linhas.append(f'{parte}: {t}' + (f' (fala: {fala})' if fala else ''))
    return '\n'.join(linhas)

wb = Workbook()
ws = wb.active
ws.title = 'Programação'
cols = [('Semana', 8), ('Data', 11), ('Dia', 10), ('Bloco', 17), ('Formato', 11), ('Título', 32), ('Subtítulo', 36),
        ('Roteiro (cena a cena ou slide a slide)', 70), ('Legenda', 70)]
for j, (nome, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=nome)
    c.font, c.fill, c.border = Font(name='Arial', size=10, bold=True, color='1F3A36'), PatternFill('solid', fgColor='96C5BD'), BORDA
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 28
ini = C['D'](19, 10)
for i, p in enumerate(C['POSTS'], 2):
    vals = [(p['data'] - ini).days // 7 + 1, p['data'], DIAS[p['data'].weekday()], p['bloco'], p['fmt'], p['titulo'], p['sub'], roteiro(p), p['legenda']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font, c.border = Font(name='Arial', size=10, color='1F3A36'), BORDA
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if j in (1, 2, 3) else None)
        if j == 2:
            c.number_format = 'DD/MM/YYYY'
    ws.cell(row=i, column=5).fill = PatternFill('solid', fgColor=COR_FMT[p['fmt']])
    ws.cell(row=i, column=6).font = Font(name='Arial', size=10, bold=True, color='1F3A36')
    n = max(sum(max(1, math.ceil(len(par) / (w * 1.15))) for par in str(v).split('\n')) for v, (_, w) in zip(vals, cols))
    ws.row_dimensions[i].height = n * 13 + 6
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:I{len(C["POSTS"]) + 1}'
wb.save(sys.argv[1])
print('ok', sys.argv[1])
