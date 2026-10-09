# Gera marketing/Calendario_Editorial_FairHealth.xlsx no formato MENSAL (um mês por vez).
# A versão anterior (90 dias) está em marketing/arquivo/ e o gerador dela em gerar_calendario_90dias.py.
# Uso: python3 marketing/ferramentas/gerar_calendario_mensal.py marketing/Calendario_Editorial_FairHealth.xlsx
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = 'Arial'
TIFFANY, TEXTO, AMARELO = '9FC5BD', '1F3A36', 'FFFF00'
thin = Side(style='thin', color='C9D6D2')
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CAB = Font(name=F, bold=True, color=TEXTO, size=10)
NORMAL = Font(name=F, size=10)
TITULO = Font(name=F, bold=True, size=14, color=TEXTO)
FILL_CAB = PatternFill('solid', fgColor=TIFFANY)
FILL_INPUT = PatternFill('solid', fgColor=AMARELO)

I, P, J = 'Institucional', 'Produto', 'Projetos, eventos e entregas'
V, C, FU = 'Vídeo', 'Carrossel', 'Foto única'
IGLI, IG = 'Instagram + LinkedIn', 'Instagram'
FJ = 'FairJob'
HT = '#fairhealth #saudecorporativa #cuidadocentradonapessoa #medicinadoestilodevida #saudedotrabalhador'
HT_NR1 = '#fairhealth #nr1 #riscospsicossociais #saudementalnotrabalho #rh'

# Argumentos do grupo (arquitetura de marcas) + a apresentação da marca
AP, LBE, NR1 = 'Apresentação', 'Liderança Bem-Estar', 'NR-1'
ARGUMENTOS = [AP, LBE, NR1, 'Stop Loss', 'Ambiente positivo', 'Marca empregadora']
# Cenas da história do mês
ABE, DOR, VIR, PRO, CON = 'Abertura', 'Dor', 'Virada', 'Prova', 'Convite'
CENAS = [ABE, DOR, VIR, PRO, CON]
# Personas (posicionamento v3.2)
RH, LID, COL, MT, RC, CFO, CL = ('RH', 'Lideranças', 'Colaboradores', 'Medicina do Trabalho',
                                 'Risco e Compliance', 'Diretoria financeira', 'RI e C-Level')
PERSONAS = [RH, LID, COL, MT, RC, CFO, CL]

ARTIGO = ('Artigo "Riscos psicossociais no ambiente de trabalho e saúde", RH Pra Você, 25/09/2026. '
          'https://rhpravoce.com.br/colab/riscos-psicossociais-no-ambiente-de-trabalho-2')
SEM_TIME = 'Sem time em cena e sem gravar nas empresas-clientes (decisão de 09/10).'
SEM_NOMES = 'Nenhum nome real de pessoa ou empresa no post (decisão de 09/10).'
NUM_SUAVE = 'Número suave: cuidado primeiro, número como contexto, fonte visível e discreta (aba Ciência).'
INSS = ('Ministério da Previdência Social: 472.328 licenças por transtornos mentais em 2024, maior número da série iniciada em 2014 '
        '(dados divulgados em 2025; conferir no painel oficial de dados abertos antes de publicar).')
DPP = ('Diabetes Prevention Program Research Group. Reduction in the incidence of type 2 diabetes with lifestyle intervention or metformin. '
       'N Engl J Med. 2002;346(6):393-403. doi:10.1056/NEJMoa012512')
OMS = ('OMS. Guidelines on mental health at work, 28/09/2022 (com a OIT: policy brief "Mental health at work"). '
       'https://www.who.int/news/item/28-09-2022-who-and-ilo-call-for-new-measures-to-tackle-mental-health-issues-at-work')
RAND = ('Hafner M. et al. Why sleep matters: the economic costs of insufficient sleep. RAND Europe, 2016 (RR-1791). '
        'https://www.rand.org/news/press/2016/11/30.html')

# deslocamento (dias a partir da segunda) das 4 publicações de cada semana; evita feriados nacionais
OFF = {1: (0, 1, 3, 4), 2: (0, 1, 3, 4), 3: (1, 2, 3, 4), 4: (0, 1, 3, 4)}   # semana 3: 02/11 Finados

ROWS = [
 # ---------------- SEMANA 1 · ABERTURA E DOR: "Isso é comigo?" ----------------
 dict(sem=1, linha=I, arg=AP, cena=ABE, tema='Apresentação da Fair Health', serie='',
  titulo='Antes de qualquer exame, uma conversa', persona=f'{COL}; {RH}',
  nec='O que é a Fair Health e o que ela faz dentro da empresa?',
  obj='Apresentar a marca e o jeito de cuidar', formato=V,
  porque='Vídeo de abertura com a pergunta da marca: apresenta em 30 s o que é, onde fica e como começa. Texto na tela para quem assiste sem som.',
  canal=IGLI, collab=FJ,
  msg='A Fair Health leva o cuidado para dentro da empresa, e todo cuidado começa com uma pergunta: o que é importante para você?',
  fonte='Posicionamento Fair Health v3.2 (promessa e pergunta).',
  valid='Roteiro e legenda citam assistentes sociais: confirmar se fazem parte do atendimento (o posicionamento cita médico e psicológico). ' + SEM_TIME + ' Definir onde gravar (local próprio ou neutro) ou montar com fotos e tipografia.',
  status='Roteiro pronto',
  texto='Antes de qualquer exame, uma conversa.\n\nEsse é o Escritório do Cuidado: um espaço dentro da empresa, com médicos, psicólogos e assistentes sociais, perto de quem trabalha.\n\nAqui, todo cuidado começa com a mesma pergunta: o que é importante para você? A partir da resposta, montamos juntos um plano que cabe na sua rotina e seguimos acompanhando.\n\nSeja bem-vindo. Nos próximos dias, a gente mostra como esse cuidado funciona.\n\n' + HT,
  roteiro='Reel de 30 a 34 s (instagram/primeiros-posts-acolhimento.md, ID 1).\n0–3 s: "Antes de qualquer exame, a gente começa com uma conversa." Porta se abrindo.\n3–9 s: o espaço (recepção, copa, sala), sem pessoas.\n9–16 s: detalhes que acolhem: mãos servindo café, cadeira puxada, sem rosto.\n16–23 s: a pergunta "O que é importante para você?" aparece na tela, em voz em off.\n23–30 s: a partir da resposta, um cuidado que cabe na rotina.\n30–34 s: logo e assinatura.',
  criativo='Vídeo 1080x1920. Texto na tela dentro da área segura. Capa no modelo M1 (Tiffany, pergunta), conferindo o recorte da grade. LinkedIn: mesmo vídeo, legenda adaptada para RH.',
  coment='Aprovado na revisão de 09/10. Ajustes pelas restrições de 09/10: a legenda não promete mais apresentar o time ("você vai conhecer quem cuida" virou "a gente mostra como esse cuidado funciona") e, no roteiro, a pergunta entra na tela em voz em off (antes cada integrante dizia um pedaço).'),
 dict(sem=1, linha=P, arg=AP, cena=ABE, tema='Escritório do Cuidado', serie='Como funciona o cuidado (1/3): o percurso',
  titulo='Aqui, o cuidado começa com uma pergunta', persona=COL,
  nec='O que acontece quando eu vou ao Escritório do Cuidado? Vou ser julgado? O que eu contar fica entre nós?',
  obj='Explicar o percurso do atendimento e reduzir a barreira de ir', formato=C,
  porque='Carrossel passo a passo: o formato mais salvo em conteúdo explicativo. Responde as objeções (julgamento, sigilo) uma por slide.',
  canal=IGLI, collab=FJ,
  msg='Primeiro a gente te ouve e entende sua rotina; depois monta o plano com você e acompanha. O que você conta fica entre você e quem cuida.',
  fonte='Cartilha do Escritório do Cuidado.',
  valid='Responsável médica confere as etapas citadas (anamnese, bioimpedância, smartband por 48 horas, retorno).',
  status='Em revisão',
  texto='Aqui, o cuidado começa com uma pergunta: o que é importante para você?\n\nAntes de exame, antes de número, a gente quer saber da sua rotina, do seu trabalho, do que anda pesando e do que você gostaria de mudar.\n\nO Escritório do Cuidado fica dentro da empresa, perto de onde você já está. Lá, médicos e psicólogos se revezam para montar com você um cuidado que caiba na sua vida, com base na Medicina do Estilo de Vida.\n\nO objetivo da fairhealth é melhorar a qualidade de vida de quem trabalha, com um cuidado feito para cada pessoa e acompanhado de perto.\n\nE o que você conta fica entre você e quem cuida.\n\nArrasta para o lado para ver como funciona.\n\n' + HT,
  roteiro='10 slides: 1 capa (pergunta) · 2 fica dentro da empresa · 3 gente que cuida de gente · 4 Passo 1: a gente te conhece · 5 Passo 2: o que é importante para você? · 6 Passo 3: entender antes de propor (bioimpedância e smartband por 48 h) · 7 Passo 4: o plano é seu · 8 Passo 5: retorno e acompanhamento · 9 sigilo · 10 convite para comentar.',
  criativo='Arte no Canva (design DAHW6JfDldU, 11 páginas): escolher entre as versões originais e as novas dos slides 1, 5, 7, 9 e 10. Modelos M1, M2 e M3. LinkedIn: exportar em PDF.',
  coment='Conferir o slide 3 ("gente que cuida de gente"): não pode ter foto ou nome de integrante do time.'),
 dict(sem=1, linha=I, arg=LBE, cena=DOR, tema='Cuidar antes (manifesto)', serie='',
  titulo='Antes do afastamento, teve um sinal', persona=f'{RH}; {LID}',
  nec='Por que as pessoas só aparecem quando já estão doentes ou afastadas?',
  obj='Plantar o fio do mês ("cuidar antes") e fazer o RH se reconhecer na dor', formato=V,
  porque='Reel de tipografia animada: não depende de filmagem nem de rosto, o gancho está na primeira frase e a cadeia do "antes" prende até o fim. Fácil de adaptar para o LinkedIn.',
  canal=IGLI, collab=FJ,
  msg='Todo afastamento tem um antes. A Fair Health cuida nesse antes, dentro da empresa, começando por uma pergunta.',
  fonte='Posicionamento Fair Health v3.2 (eixo "antes"). ' + INSS,
  valid='Sem promessa de evitar doença (CFM): o texto fala de cuidado, não de resultado. ' + NUM_SUAVE + ' Conferir o número no painel oficial.',
  status='Pauta proposta',
  texto='Antes do afastamento, teve um atestado.\nAntes do atestado, teve uma semana pesada.\nAntes da semana pesada, teve uma noite sem dormir.\n\nE antes de tudo isso, ninguém perguntou: o que é importante para você?\n\nEm 2024, o Brasil registrou mais de 470 mil licenças do trabalho por saúde mental, o maior número em dez anos (Ministério da Previdência Social). Cada uma delas teve um antes.\n\nA Fair Health existe para cuidar nesse antes. Dentro da empresa, perto de quem trabalha, com médicos e psicólogos e um plano feito para cada pessoa.\n\nCuidar antes também é papel da liderança. Na sua empresa, quem faz essa pergunta?\n\n' + HT,
  roteiro='Reel de 22 a 26 s, só tipografia sobre o fundo Tiffany, frase a frase, com corte seco.\n0–3 s: "Antes do afastamento, teve um atestado."\n3–6 s: "Antes do atestado, teve uma semana pesada."\n6–9 s: "Antes da semana pesada, uma noite sem dormir."\n9–14 s: "E antes de tudo isso, ninguém perguntou:"\n14–18 s: "O que é importante para você?" (palavra-chave em destaque)\n18–22 s: em texto menor e calmo: "Em 2024, mais de 470 mil licenças por saúde mental no Brasil. Cada uma teve um antes." Fonte no rodapé.\n22–26 s: "Cuidar antes." + logo.\nTrilha calma, sem locução (ou voz em off opcional).',
  criativo='Modelo M9 (vídeo) só com tipografia Raleway. Capa: "Antes do afastamento, teve um sinal". Sem imagem de banco.'),
 dict(sem=1, linha=J, arg=NR1, cena=DOR, tema='Na imprensa: artigo sobre riscos psicossociais', serie='Saúde mental é de todos (1/2)',
  titulo='O trabalho não é o único vilão', persona=f'{RH}; {LID}; {RC}',
  nec='A NR-1 colocou os riscos psicossociais na conta da empresa. O que é responsabilidade da empresa e o que não é?',
  obj='Trazer a NR-1 com uma visão madura e transferir a autoridade da publicação para a marca', formato=C,
  porque='Carrossel-resumo de artigo publicado: formato muito compartilhado entre profissionais de RH. Leva a autoridade do veículo para a marca, sem citar pessoas.',
  canal=IGLI, collab=f'{FJ}; RH Pra Você (marcar o veículo, se fizer sentido)',
  msg='A NR-1 formalizou que o trabalho pode adoecer, e a empresa responde pelo ambiente que constrói. Mas saúde mental é multifatorial: é, ao mesmo tempo, de cada um e de todos.',
  fonte=ARTIGO + ' Dado do slide 4: Global Burden of Disease (citado no artigo).',
  valid='Citar só o artigo (título e veículo), sem o nome do autor. ' + SEM_NOMES + ' Conferir as citações literais com o artigo. ' + NUM_SUAVE,
  status='Pauta proposta',
  texto='O trabalho pode adoecer. Mas não é o único vilão.\n\nDesde maio de 2026, a NR-1 exige que as empresas identifiquem, avaliem e gerenciem os riscos psicossociais. É um avanço necessário: assédio, excesso de pressão, jornadas inadequadas e lideranças despreparadas têm impacto real na saúde mental.\n\nUm artigo publicado na RH Pra Você propõe ampliar o debate: ninguém começa a vida quando é contratado. Família, escola, comunidade e condições de vida chegam ao trabalho junto com a pessoa.\n\nIsso não diminui a responsabilidade da empresa. Ela responde pelo ambiente que constrói. Mas não controla todas as variáveis.\n\n"Saúde mental é, ao mesmo tempo, de cada um e de todos."\n\nO artigo completo está no link dos stories. Concorda? Conta pra gente nos comentários.\n\n' + HT_NR1,
  roteiro='8 slides.\n1 Capa: "O trabalho pode adoecer. Mas não é o único vilão." + "Da RH Pra Você".\n2 Desde maio de 2026, a NR-1 exige identificar, avaliar e gerenciar os riscos psicossociais.\n3 É um avanço necessário: ambientes tóxicos, assédio, excesso de pressão, jornadas inadequadas, insegurança e lideranças despreparadas adoecem.\n4 "Nenhuma pessoa começa sua vida quando é contratada por uma empresa." Abaixo, em texto menor: no mundo, quase 1 bilhão de pessoas convivem com algum transtorno mental (Global Burden of Disease, 2019).\n5 A empresa responde pelo ambiente que constrói. Não controla todas as variáveis.\n6 Responsabilidade compartilhada: poder público, sociedade, empresas e cada pessoa (sem culpar quem adoece).\n7 "Saúde mental é, ao mesmo tempo, de cada um e de todos." (do artigo, sem nome)\n8 A parte da empresa começa antes: um cuidado dentro do trabalho, que começa com uma pergunta. Link nos stories.',
  criativo='Modelo M8 (citação) nos slides 4 e 7 e M2 nos demais. Sem foto. Stories no mesmo dia com o link do artigo.'),

 # ---------------- SEMANA 2 · VIRADA: "Como a Fair Health resolve?" ----------------
 dict(sem=2, linha=P, arg=LBE, cena=VIR, tema='Medicina do Estilo de Vida', serie='Pilares da MEV (abertura)',
  titulo='Medicina do Estilo de Vida em 1 minuto', persona=f'{COL}; {RH}',
  nec='O que é Medicina do Estilo de Vida? É dieta? É academia?',
  obj='Apresentar a base científica do cuidado e abrir a série dos pilares', formato=V,
  porque='Explicador animado com ícones (motion): conceito novo explicado em 1 minuto, sem rosto. Responde ao mito "é só dieta e academia".',
  canal=IGLI, collab=FJ,
  msg='Seis pilares, uma pessoa inteira: alimentação, atividade física, sono, manejo do estresse, relacionamentos e evitar substâncias de risco. O plano olha para todos, a partir da rotina real.',
  fonte='ACLM (American College of Lifestyle Medicine), 6 pilares; CBMEV (Colégio Brasileiro de Medicina do Estilo de Vida); cartilha do Escritório do Cuidado.',
  valid='6 pilares, sem espiritualidade (decisão de 09/10, ainda a confirmar). Responsável médica valida a definição de MEV e os nomes dos pilares.',
  status='Pauta proposta',
  roteiro='Reel de 45 a 60 s, motion com ícones, voz em off.\n0–3 s: "Medicina do Estilo de Vida é dieta e academia? Não só."\n3–10 s: definição curta, com fonte na tela.\n10–40 s: os 6 pilares, um ícone e uma frase por pilar: alimentação · atividade física · sono · manejo do estresse · relacionamentos · evitar substâncias de risco.\n40–50 s: "Os pilares conversam entre si: o plano olha para a pessoa inteira."\n50–60 s: "No Escritório do Cuidado, começa com uma pergunta." + logo.',
  criativo='Modelo M9 com ícones lineares em verde-escuro sobre Tiffany. Etiqueta "PILARES".'),
 dict(sem=2, linha=I, arg=LBE, cena=VIR, tema='Tese da marca', serie='',
  titulo='Não é benefício. É infraestrutura de cuidado.', persona=f'{RH}; {CL}',
  nec='Isso é mais um benefício para a lista? Por que seria diferente do que já temos?',
  obj='Fixar a frase-tese e posicionar a categoria sem citar concorrentes', formato=C,
  porque='Carrossel comparativo de duas colunas (isto × aquilo): a diferença fica clara em segundos e o post é salvo e repassado ao decisor.',
  canal=IGLI, collab=FJ,
  msg='Benefício é algo que a pessoa usa se lembrar. Infraestrutura de cuidado está dentro da empresa todo dia, começa antes do problema e tem gente acompanhando.',
  fonte='Posicionamento Fair Health v3.2 (diferenciais).',
  valid='Não citar marcas de terceiros nem operadoras. Sem números.',
  status='Pauta proposta',
  roteiro='6 slides.\n1 Capa: "Não é benefício. É infraestrutura de cuidado."\n2 Benefício: a pessoa usa se lembrar. Infraestrutura: está lá todo dia, dentro da empresa.\n3 Benefício: espera o problema aparecer. Infraestrutura: começa antes.\n4 Benefício: é um cartão. Infraestrutura: é gente acompanhando, com um plano para cada pessoa.\n5 Benefício: termina na adesão. Infraestrutura: o RH acompanha o cuidado por indicadores do grupo.\n6 "O que é importante para você?" Comente CUIDADO para saber como funciona.',
  criativo='Modelo M1 na capa e comparativo em duas colunas (Tiffany × verde-escuro) nos slides 2 a 5.'),
 dict(sem=2, linha=J, arg=NR1, cena=VIR, tema='Entregas para o RH', serie='Tecnologia a favor da pessoa',
  titulo='O que o RH vê (e o que não vê)', persona=f'{RH}; {RC}',
  nec='O que eu recebo para acompanhar o programa? Vou ver dados individuais?',
  obj='Mostrar a entrega para a empresa e a regra de privacidade', formato=V,
  porque='Gravação de tela de um painel genérico, com texto na tela: mostra o produto de verdade sem expor ninguém e responde a dúvida que trava a decisão do RH.',
  canal=IGLI, collab=f'{FJ}; fairdata',
  msg='O RH acompanha o cuidado por indicadores do grupo, anonimizados e integrados à fairdata. O que cada pessoa conta fica no atendimento.',
  fonte='Posicionamento Fair Health v3.2 (fairdata).',
  valid='Time confirma quais indicadores entram no relatório. Painel com dados fictícios e aviso "dados ilustrativos". Não prometer conformidade com a NR-1.',
  status='Pauta proposta',
  roteiro='Reel de 25 a 30 s, gravação de tela.\n0–3 s: "O que o RH vê no Escritório do Cuidado?"\n3–15 s: rolagem pelo painel (adesão, retorno, temas mais procurados), dados ilustrativos.\n15–22 s: "O que o RH não vê: o que cada pessoa contou."\n22–30 s: "Indicadores do grupo para cuidar antes." + logo.',
  criativo='Modelo M9. Painel montado com dados fictícios, moldura de navegador ou de notebook.'),
 dict(sem=2, linha=P, arg=AP, cena=VIR, tema='Plano Individualizado de Cuidado', serie='Como funciona o cuidado (2/3): o plano',
  titulo='O plano é seu: como nasce o Plano Individualizado de Cuidado', persona=COL,
  nec='Vou receber uma lista de regras impossíveis de seguir?',
  obj='Explicar como o plano é montado com a pessoa (decisão compartilhada)', formato=C,
  porque='Carrossel com card de interface (M3), já testado no carrossel de acolhimento, contado pela história de uma personagem fictícia: história prende mais que explicação e mostra o plano como algo possível.',
  canal=IGLI, collab=FJ,
  msg='O plano parte do que é importante para você e da sua rotina real, com ações nos pilares da MEV, e muda quando a vida muda.',
  fonte='Cartilha do Escritório do Cuidado. Personagem fictícia.',
  valid='Responsável médica valida etapas e prazos. Nome único do plano: "Plano Individualizado de Cuidado" ou "Plano de Cuidado Personalizado"? Indicar no slide e na legenda que é uma história ilustrativa (personagem fictícia), para não parecer depoimento.',
  status='Pauta proposta',
  roteiro='7 slides, com a Clara (personagem fictícia, 41 anos, supervisora de turno).\n1 Capa: "O plano é seu."\n2 A Clara chegou dizendo: "Quero voltar a ter energia para brincar com meu filho." Isso é o que importa para ela.\n3 Card do plano da Clara: 3 ações pequenas (dormir 30 min mais cedo, caminhada de 10 min no almoço, conversa quinzenal com a psicóloga).\n4 Cabe na rotina real: no turno dela, não numa rotina ideal.\n5 Muda quando a vida muda.\n6 Alguém acompanha, sem julgamento.\n7 "E para você, o que é importante?" Rodapé: "História ilustrativa. Personagem fictícia."',
  criativo='Modelo M3 (card de interface do plano). Sem foto de pessoa: a personagem aparece só no texto e no card.'),

 # ---------------- SEMANA 3 · PROVA: "Por que acreditar?" ----------------
 dict(sem=3, linha=J, arg=LBE, cena=PRO, tema='Adesão (número liberado)', serie='',
  titulo='Quando o cuidado está perto, as pessoas voltam', persona=f'{RH}; {CFO}',
  nec='Programa de saúde que ninguém usa é custo. As pessoas usam isso de verdade?',
  obj='Provar adesão com número liberado, contado como cuidado e não como resultado', formato=V,
  porque='Número contado como história: primeiro a cena de cuidado, depois o número, em tom calmo. É a prova que o RH leva para a diretoria, sem cara de anúncio.',
  canal=IGLI, collab=FJ,
  msg='Quando o cuidado está perto e começa por uma conversa, as pessoas voltam: numa indústria atendida, 8 em cada 10 pessoas voltaram para mais de um atendimento.',
  fonte='Dados internos de um cliente do setor industrial, 1º semestre de 2026 (liberados; ver aba Fontes). Sempre agregado e sem o nome da empresa.',
  valid=SEM_NOMES + ' Usar "uma indústria". ' + NUM_SUAVE,
  status='Pauta proposta',
  roteiro='Reel de 18 a 22 s, tipografia e números suaves.\n0–4 s: "Ninguém volta a um lugar onde não se sentiu cuidado."\n4–9 s: "Numa indústria que atendemos, no primeiro semestre de 2026..."\n9–14 s: "8 em cada 10 pessoas voltaram para mais de um atendimento." (número em destaque, sem exagero)\n14–18 s: em texto menor: "A adesão chegou a 98%."\n18–22 s: "Cuidado que as pessoas procuram." + fonte discreta + logo.',
  criativo='Modelo M6 (dado com fonte) animado, em ritmo lento. Número em verde-escuro, nunca em vermelho; sem setas para cima nem efeito de "conquista".'),
 dict(sem=3, linha=P, arg=AP, cena=PRO, tema='Tecnologia: wearables', serie='Como funciona o cuidado (3/3): o acompanhamento',
  titulo='O cuidado não para quando a consulta acaba', persona=f'{COL}; {RH}',
  nec='Para que serve a pulseira? Vão me vigiar?',
  obj='Explicar o acompanhamento com tecnologia e a regra de privacidade', formato=C,
  porque='Carrossel explicativo com a objeção no título ("vão me vigiar?"): responde a dúvida mais comum antes que ela vire resistência.',
  canal=IGLI, collab=FJ,
  msg='A smartband e a bioimpedância ajudam a olhar, junto com a pessoa, como foram as últimas semanas. Os dados ficam no atendimento; o RH vê só o grupo.',
  fonte='Cartilha do Escritório do Cuidado.',
  valid='[PENDENTE] Quais dados a smartband registra e se todos os atendidos usam a pulseira (instagram/posts/wearables/README.md).',
  status='Em revisão',
  criativo='Arte pronta em instagram/posts/wearables (8 slides). Capa: foto de pulso sem rosto (modelo com autorização, fora das empresas-clientes) ou capa só tipográfica.'),
 dict(sem=3, linha=I, arg=AP, cena=PRO, tema='Origem', serie='',
  titulo='Tudo começou num hospital', persona=f'{RH}; {MT}',
  nec='De onde vem esse modelo? É só mais um programa de bem-estar?',
  obj='Dar lastro à marca com a história do modelo', formato=V,
  porque='Linha do tempo animada (ano grande + uma frase): história de origem gera confiança e não precisa mostrar pessoas.',
  canal=IGLI, collab=FJ,
  msg='O modelo nasceu em 2009, num hospital, para colocar o paciente no centro do cuidado. Hoje, coloca quem trabalha no centro, dentro da empresa.',
  fonte='Histórico do Escritório do Paciente (2009); caso do hospital onde o modelo começou (liberado; ver aba Fontes).',
  valid='Sem nomes de pessoas nem do hospital. Não usar "validado no SUS" até confirmar a fonte. Conferir datas.',
  status='Pauta proposta',
  roteiro='Reel de 25 a 30 s, linha do tempo em tipografia.\n0–3 s: "Tudo começou num hospital."\n3–10 s: 2009: Escritório do Paciente, o paciente no centro.\n10–18 s: o que o hospital aprendeu: ouvir antes de propor.\n18–25 s: hoje: o Escritório do Cuidado, dentro da empresa.\n25–30 s: "Mesma pergunta: o que é importante para você?" + logo.\nAlternativa: carrossel no modelo M2.',
  criativo='Modelo M9 com marcos da linha do tempo (modelo M2 animado). Sem fotos de pessoas.'),
 dict(sem=3, linha=J, arg=NR1, cena=PRO, tema='Na imprensa: artigo sobre riscos psicossociais', serie='Saúde mental é de todos (2/2)',
  titulo='A pergunta maior', persona=f'{RH}; {LID}',
  nec='Como olhar para a saúde mental do time sem simplificar?',
  obj='Reaproveitar o artigo com uma frase forte e reforçar a autoridade', formato=FU,
  porque='Card de citação: uma frase forte, fácil de compartilhar nos stories e no LinkedIn. Reaproveita o artigo sem repetir o carrossel.',
  canal=IGLI, collab=FJ,
  msg='"Antes de perguntar apenas o que o trabalho está fazendo com a saúde mental das pessoas, deveríamos fazer uma pergunta maior: o que a nossa forma de viver em sociedade está fazendo com a saúde mental das pessoas?"',
  fonte=ARTIGO,
  valid='Citação literal conferida com o artigo. Crédito só ao artigo e ao veículo, sem nome do autor. ' + SEM_NOMES,
  status='Pauta proposta',
  roteiro='Imagem única: a pergunta em destaque, assinatura "Do artigo \'Riscos psicossociais no ambiente de trabalho e saúde\', RH Pra Você". Legenda curta que liga a pergunta à da marca ("o que é importante para você?") e convida a ler o artigo.',
  criativo='Modelo M8 (citação), sem foto.'),

 # ---------------- SEMANA 4 · CONVITE: "Como começo?" ----------------
 dict(sem=4, linha=P, arg=LBE, cena=CON, tema='Pilar: sono', serie='Pilares da MEV (1: sono)',
  titulo='Dormir não é tempo perdido', persona=COL,
  nec='Por que eu acordo cansado? O que dá para mudar na rotina?',
  obj='Educar sobre o pilar sono e convidar o colaborador a procurar o Escritório', formato=V,
  porque='Reel de "3 hábitos" com texto na tela sobre imagens de banco: dica prática e salvável, sem rosto e sem gravar nas empresas.',
  canal=IG, collab=FJ,
  msg='Dormir não é tempo perdido. A faixa saudável para adultos fica entre 7 e 9 horas por noite. No Escritório do Cuidado, o sono entra no plano a partir da sua rotina real.',
  fonte='Guia "Saúde mental: o cuidado no dia a dia" (higiene do sono; OMS, APA, IPq-USP). ' + RAND,
  valid='Responsável médica valida as orientações. Sem promessa clínica. ' + NUM_SUAVE,
  status='Pauta proposta',
  roteiro='Reel de 22 a 27 s.\n0–3 s: "Acorda cansado mesmo dormindo?"\n3–6 s: em texto calmo: "Para adultos, o sono saudável fica entre 7 e 9 horas." (fonte discreta: RAND Europe, 2016)\n6–19 s: 3 hábitos, um por cena (tela antes de dormir, horário regular, cafeína à tarde).\n19–27 s: "Se o cansaço não passa, conversa com a gente no Escritório do Cuidado." + logo.\nLegenda (opcional, para o RH no LinkedIn): um estudo da RAND Europe estimou que a falta de sono custa mais de 1 milhão de dias de trabalho por ano só nos Estados Unidos.',
  criativo='Modelo M9 com etiqueta "PILARES · SONO". Imagens de banco (quarto, celular à noite), sem rosto identificável.'),
 dict(sem=4, linha=I, arg=LBE, cena=CON, tema='A pergunta (com a recomendação da OMS)', serie='Ciência no cuidado',
  titulo='Uma pergunta que todo líder pode fazer', persona=f'{LID}; {RH}',
  nec='Como a liderança pode cuidar do time no dia a dia, sem ser médica nem psicóloga?',
  obj='Levar a pergunta da marca para a liderança, com o respaldo da recomendação da OMS, e gerar conversa', formato=C,
  porque='Carrossel de ciência aplicada: a recomendação de uma fonte respeitada (OMS) vira um gesto simples que o líder pode repetir. Fecha com pergunta aberta, que gera comentários.',
  canal=IGLI, collab=FJ,
  msg='A OMS recomenda preparar as lideranças para prevenir ambientes de estresse e acolher quem está em sofrimento. Um bom começo cabe numa pergunta: o que é importante para você?',
  fonte=OMS,
  valid='Conferir a redação da recomendação nas diretrizes da OMS. O líder não substitui o cuidado clínico: deixar claro onde pedir ajuda.',
  status='Pauta proposta',
  roteiro='6 slides.\n1 Capa: "Uma pergunta que todo líder pode fazer."\n2 Em 2022, a OMS publicou suas primeiras diretrizes sobre saúde mental no trabalho. Uma das recomendações: preparar as lideranças.\n3 Para quê? Para prevenir ambientes de estresse e acolher quem está em sofrimento.\n4 Não é virar terapeuta. É perguntar, ouvir e saber para onde encaminhar.\n5 Um bom começo: "O que é importante para você?"\n6 Líder, quando foi a última vez que você fez essa pergunta? Conta nos comentários. Fonte: OMS, 2022.\nStories no mesmo dia com caixa de perguntas.',
  criativo='Modelo M1 na capa e M2 nos demais; fonte no rodapé (M6).'),
 dict(sem=4, linha=J, arg=AP, cena=CON, tema='Como começa numa empresa', serie='Como chega a uma empresa',
  titulo='Como o Escritório do Cuidado chega à sua empresa', persona=f'{RH}; {CFO}',
  nec='Quero isso na minha empresa. Por onde começa?',
  obj='Deixar clara a oferta e o primeiro passo comercial', formato=V,
  porque='Passo a passo animado com chamada para ação por palavra-chave ("comente CUIDADO"): fecha o mês com um caminho claro para o RH.',
  canal=IGLI, collab=FJ,
  msg='Começa com um diagnóstico gratuito, segue com uma proposta personalizada e chega à implantação. Comente CUIDADO.',
  fonte='Executive Summary (etapas comerciais).',
  valid='Não citar prazos de cada etapa sem confirmação. Confirmar o destino do link (www.fairhealth.com.br) e quem responde os comentários e DMs.',
  status='Pauta proposta',
  roteiro='Reel de 20 a 25 s, motion em 3 passos.\n0–3 s: "Como levar o Escritório do Cuidado para a sua empresa?"\n3–9 s: PASSO 1 · Diagnóstico gratuito\n9–15 s: PASSO 2 · Proposta personalizada\n15–20 s: PASSO 3 · Implantação e primeiros atendimentos\n20–25 s: "Comente CUIDADO." + logo.',
  criativo='Modelo M2 (passos) animado e M1 no fim.'),
 dict(sem=4, linha=P, arg=LBE, cena=CON, tema='Ciência: estilo de vida e diabetes (Dia Mundial do Diabetes, 14/11)', serie='Ciência no cuidado',
  titulo='Estilo de vida também é tratamento sério', persona=f'{COL}; {RH}; {MT}',
  nec='Mudar hábitos faz diferença de verdade ou é só conselho?',
  obj='Mostrar a base científica da MEV com um estudo clássico, na véspera do Dia Mundial do Diabetes', formato=C,
  porque='Carrossel "o que a ciência diz": um estudo, uma ideia, a fonte completa. Conecta uma data do calendário ao nosso trabalho com evidência, sem oportunismo.',
  canal=IGLI, collab=FJ,
  msg='Num estudo clássico, publicado no New England Journal of Medicine, mudanças de estilo de vida reduziram em 58% os novos casos de diabetes tipo 2 em pessoas com risco, em comparação com o placebo. É por isso que o nosso cuidado começa pela rotina.',
  fonte=DPP,
  valid='Responsável médica valida a leitura do estudo. Deixar claro que o estudo foi com adultos com pré-diabetes (risco aumentado) e que não é promessa individual. ' + NUM_SUAVE,
  status='Pauta proposta',
  roteiro='7 slides.\n1 Capa: "Estilo de vida também é tratamento sério."\n2 Amanhã é o Dia Mundial do Diabetes. Vale lembrar de um estudo que mudou a forma de cuidar.\n3 Mais de 3 mil adultos com risco aumentado de diabetes tipo 2 foram acompanhados por quase 3 anos.\n4 Quem mudou a rotina (alimentação e 150 minutos de atividade física por semana) teve 58% menos casos novos do que o grupo placebo.\n5 Em texto menor: mais do que o grupo que recebeu medicamento (31%).\n6 Não é promessa: é a ciência dizendo que a rotina importa. Por isso o nosso cuidado começa nela, com um plano feito para cada pessoa.\n7 "O que é importante para você?" Fonte completa: Diabetes Prevention Program, N Engl J Med, 2002.',
  criativo='Modelo M6 (dado com fonte) suave nos slides 4 e 5, M2 nos demais. Números em verde-escuro, sem destaque alarmista.'),
]

# checagens: 16 no mês, 4 por semana, sem linha repetida em sequência, vídeo intercalado
assert len(ROWS) == 16
for k in range(1, len(ROWS)):
    assert ROWS[k]['linha'] != ROWS[k - 1]['linha'], (k, ROWS[k]['titulo'])
for k, r in enumerate(ROWS):
    assert (r['formato'] == V) == (k % 2 == 0), (k, r['titulo'], r['formato'])
    assert r['sem'] == k // 4 + 1, (k, r['titulo'])

COLS = [
    ('ID', 8, None), ('Mês', 8, None), ('Semana', 8, None), ('Data sugerida', 12, None),
    ('Linha editorial', 16, 'linha'), ('Argumento', 16, 'arg'), ('Cena', 11, 'cena'), ('Tema', 20, 'tema'),
    ('Série (parte)', 22, 'serie'), ('Título provisório', 36, 'titulo'), ('Persona', 20, 'persona'),
    ('Necessidade ou pergunta da persona', 34, 'nec'), ('Objetivo', 28, 'obj'), ('Formato', 11, 'formato'),
    ('Por que este formato', 40, 'porque'), ('Canal', 18, 'canal'), ('Collab / perfis', 24, 'collab'),
    ('Mensagem principal', 44, 'msg'), ('Fonte', 34, 'fonte'), ('Validação necessária', 40, 'valid'),
    ('Status', 16, 'status'), ('Texto / legenda', 50, 'texto'), ('Roteiro / estrutura', 50, 'roteiro'),
    ('Criativo (modelo e arte)', 40, 'criativo'), ('Responsável', 14, None), ('Comentários da revisão', 36, 'coment'),
]
COLUNA = {nome: get_column_letter(j) for j, (nome, _, _) in enumerate(COLS, 1)}
CL_LINHA, CL_ARG, CL_CENA, CL_FMT, CL_VAL, CL_STATUS, CL_MES = (
    COLUNA['Linha editorial'], COLUNA['Argumento'], COLUNA['Cena'], COLUNA['Formato'], COLUNA['Validação necessária'],
    COLUNA['Status'], COLUNA['Mês'])

wb = Workbook()

# ---------- Como usar ----------
ws0 = wb.active
ws0.title = 'Como usar'
linhas = [
    ('Calendário editorial · Fair Health · Instagram e LinkedIn', TITULO),
    ('Versão de 09/10/2026: planejamento MENSAL (um mês por vez). A versão anterior, de 90 dias, está em marketing/arquivo/.', NORMAL),
    ('', NORMAL),
    ('O mês conta uma história', CAB),
    ('Fio de todos os meses: "Cuidar antes". O plano de saúde trata depois; a Fair Health cuida antes, dentro da empresa, começando por uma pergunta.', NORMAL),
    ('Cada mês trabalha um argumento do grupo (aba Visão do mês) e cada semana é uma cena: Abertura/Dor ("isso é comigo?"), Virada ("como a Fair Health resolve?"), Prova ("por que acreditar?") e Convite ("como começo?").', NORMAL),
    ('Base: instagram/posicionamento-e-plano-mensal.md (análise da arquitetura de marcas e do posicionamento v3.2).', NORMAL),
    ('', NORMAL),
    ('Restrições em vigor (09/10)', CAB),
    ('1. Não apresentamos o time por enquanto (questão administrativa). A série Quem cuida está na aba Banco de ideias, com status Adiada.', NORMAL),
    ('2. Não filmamos clientes: nada gravado dentro das empresas-clientes nem com pessoas atendidas. Os vídeos do mês são de tipografia, motion, gravação de tela ou imagem de banco (aba Formatos).', NORMAL),
    ('3. Nunca citamos nomes reais de pessoas ou empresas nos posts. Quando a história precisa de alguém, usamos personagem fictícia, sinalizada como "história ilustrativa". O artigo da RH Pra Você (25/09/2026) é citado só pelo título e pelo veículo.', NORMAL),
    ('4. Números são bem-vindos, mas o foco é o cuidado: a frase de cuidado vem primeiro e o número entra como contexto, em tom calmo, com a fonte visível (regras na aba Ciência).', NORMAL),
    ('5. Artigos científicos e fontes oficiais viram posts que conectam a evidência ao nosso trabalho (aba Ciência). Medicina do Estilo de Vida: 6 pilares, sem espiritualidade (a confirmar).', NORMAL),
    ('', NORMAL),
    ('Como trabalhamos nesta planilha', CAB),
    ('1. Revisão mensal (conjunta): antes de o mês começar, revisamos temas, títulos, personas e datas. Pauta aprovada muda de "Pauta proposta" para "Pauta aprovada". Ao fim do mês, montamos o próximo a partir do Banco de ideias.', NORMAL),
    ('2. Detalhamento semanal: a cada semana, detalhamos texto, roteiro e criativo das peças da semana seguinte e passamos pela validação. A semana 1 já está detalhada.', NORMAL),
    ('3. Comentários: use a coluna "Comentários da revisão" ou comentários de célula. Não apague pautas descartadas: mude o status para "Descartada".', NORMAL),
    ('4. Datas: a data de início (célula amarela na aba Calendário) recalcula todas as datas sugeridas. Para mudar uma data só, digite a nova data por cima.', NORMAL),
    ('', NORMAL),
    ('Cadência proposta (a validar)', CAB),
    ('4 publicações por semana no feed, alternando vídeo e estático/carrossel (2 vídeos por semana). As três linhas editoriais entram em rodízio: duas peças seguidas nunca são da mesma linha.', NORMAL),
    ('Os dias da semana NÃO estão definidos: as datas são uma distribuição sugerida (evitando feriados nacionais). Stories ficam fora da contagem e servem de apoio: link do artigo, caixa de perguntas, repost.', NORMAL),
    ('', NORMAL),
    ('Marcações', CAB),
    ('[PENDENTE] = informação que ainda precisa ser confirmada. [SUGESTÃO] = posição ou texto proposto que precisa de aprovação.', NORMAL),
    ('Coluna "Validação necessária": tudo o que precisa de aval antes de publicar. Afirmações clínicas, benefícios, prazos e resultados passam pela responsável médica.', NORMAL),
    ('Usamos números liberados (aba Fontes) e números de estudos científicos e fontes oficiais com a referência completa (aba Ciência). Números do posicionamento sem fonte (23%, R$ 2.400, burnout) ficam de fora.', NORMAL),
    ('', NORMAL),
    ('Fluxo de status', CAB),
    ('Ideia → Pauta proposta → Pauta aprovada → Aguardando material → Roteiro pronto → Em produção → Em revisão → Validação médica → Aprovado → Agendado → Publicado (ou Descartada / Aguardando pauta / Adiada)', NORMAL),
    ('', NORMAL),
    ('Abas', CAB),
    ('Visão do mês: argumento, a história semana a semana e o equilíbrio do mês (contagens automáticas).', NORMAL),
    ('Calendário: as 16 pautas do mês, uma por linha, com Argumento, Cena, Persona e "Por que este formato".', NORMAL),
    ('Formatos: os formatos usados no mês, por que funcionam e o que cada um precisa para ser produzido sem time em cena e sem gravar nas empresas.', NORMAL),
    ('Ciência: estudos e fontes oficiais com referência, o que encontraram, como apresentar o número com suavidade e a ligação com o nosso trabalho.', NORMAL),
    ('Fontes: o artigo da RH Pra Você, os números internos liberados e o que não usar.', NORMAL),
    ('Banco de ideias: o que ficou para depois (com o motivo) e os temas dos próximos meses.', NORMAL),
    ('Linhas e personas · Pendências · Listas (valores das listas suspensas).', NORMAL),
]
for i, (t, f) in enumerate(linhas, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws0.column_dimensions['A'].width = 140

# ---------- Calendário ----------
ws = wb.create_sheet('Calendário')
ws['A1'] = 'Calendário editorial · Fair Health · Mês 1'
ws['A1'].font = TITULO
ws['A2'] = 'Início do mês (segunda-feira):'
ws['A2'].font = CAB
ws.merge_cells('A2:C2')
ws['D2'] = dt.date(2026, 10, 19)
ws['D2'].number_format = 'DD/MM/YYYY'
ws['D2'].fill = FILL_INPUT
ws['D2'].font = Font(name=F, bold=True, color='0000FF', size=10)
ws['D2'].comment = Comment('Data proposta para começar o mês (sugestão da revisão de 08/10). Mude aqui e todas as datas sugeridas se ajustam.', 'Planejamento')
ws['E2'] = 'Mude a data amarela para reposicionar o mês. Os dias da semana são só sugestão.'
ws['E2'].font = Font(name=F, italic=True, size=9)
HR = 4
for j, (nome, larg, _) in enumerate(COLS, 1):
    c = ws.cell(row=HR, column=j, value=nome)
    c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(j)].width = larg
ws.row_dimensions[HR].height = 30

for k, r in enumerate(ROWS):
    row = HR + 1 + k
    pos = k % 4
    ws.cell(row=row, column=1, value=f'M1-{k + 1:02d}')
    ws.cell(row=row, column=2, value='Mês 1')
    ws.cell(row=row, column=3, value=r['sem'])
    d = ws.cell(row=row, column=4, value=f'=$D$2+(C{row}-1)*7+{OFF[r["sem"]][pos]}')
    d.number_format = 'DD/MM/YYYY'
    for j, (_, _, key) in enumerate(COLS, 1):
        if key:
            ws.cell(row=row, column=j, value=r.get(key, '') or None)
    for j in range(1, len(COLS) + 1):
        c = ws.cell(row=row, column=j)
        c.font, c.alignment, c.border = NORMAL, WRAP, BORDA
last = HR + len(ROWS)
ultima_col = get_column_letter(len(COLS))
ws.freeze_panes = f'{COLUNA["Persona"]}5'
ws.auto_filter.ref = f'A{HR}:{ultima_col}{last + 60}'

for valor, cor in ((I, 'D7EAE5'), (P, 'F7EBC5'), (J, 'DFE4F0')):
    ws.conditional_formatting.add(f'{CL_LINHA}{HR + 1}:{CL_LINHA}{last + 60}',
        FormulaRule(formula=[f'${CL_LINHA}{HR + 1}="{valor}"'], fill=PatternFill('solid', fgColor=cor)))
ws.conditional_formatting.add(f'{CL_FMT}{HR + 1}:{CL_FMT}{last + 60}',
    FormulaRule(formula=[f'${CL_FMT}{HR + 1}="{V}"'], font=Font(name=F, bold=True, color='1F3A36')))
for valor, cor in ((ABE, 'EEF5F3'), (DOR, 'F6E3DF'), (VIR, 'E3EEF6'), (PRO, 'E6F2E3'), (CON, 'F3ECF7')):
    ws.conditional_formatting.add(f'{CL_CENA}{HR + 1}:{CL_CENA}{last + 60}',
        FormulaRule(formula=[f'${CL_CENA}{HR + 1}="{valor}"'], fill=PatternFill('solid', fgColor=cor)))

# ---------- Listas ----------
wl = wb.create_sheet('Listas')
STATUS = ['Ideia', 'Pauta proposta', 'Pauta aprovada', 'Aguardando material', 'Aguardando pauta', 'Roteiro pronto',
          'Em produção', 'Em revisão', 'Validação médica', 'Aprovado', 'Agendado', 'Publicado', 'Descartada', 'Adiada']
LISTAS = {
    'A': ('Linha editorial', [I, P, J]),
    'B': ('Formato', [V, C, FU, 'Stories']),
    'C': ('Canal', [IGLI, IG, 'Instagram Stories', 'LinkedIn']),
    'D': ('Status', STATUS),
    'E': ('Mês', [f'Mês {n}' for n in range(1, 13)]),
    'F': ('Argumento', ARGUMENTOS),
    'G': ('Cena', CENAS),
    'H': ('Persona', PERSONAS),
}
for col, (tit, vals) in LISTAS.items():
    wl[f'{col}1'] = tit
    wl[f'{col}1'].font, wl[f'{col}1'].fill = CAB, FILL_CAB
    for i, v in enumerate(vals, 2):
        wl[f'{col}{i}'] = v
        wl[f'{col}{i}'].font = NORMAL
    wl.column_dimensions[col].width = 26

def dv(ws_, col_lista, n, alvo):
    v = DataValidation(type='list', formula1=f"=Listas!${col_lista}$2:${col_lista}${n + 1}", allow_blank=True)
    v.error = 'Escolha um valor da lista (aba Listas). Para criar um valor novo, inclua-o na aba Listas.'
    v.errorStyle = 'warning'
    ws_.add_data_validation(v)
    v.add(alvo)
fim = last + 60
dv(ws, 'A', 3, f'{CL_LINHA}{HR + 1}:{CL_LINHA}{fim}')
dv(ws, 'B', 4, f'{CL_FMT}{HR + 1}:{CL_FMT}{fim}')
dv(ws, 'C', 4, f'{COLUNA["Canal"]}{HR + 1}:{COLUNA["Canal"]}{fim}')
dv(ws, 'D', len(STATUS), f'{CL_STATUS}{HR + 1}:{CL_STATUS}{fim}')
dv(ws, 'E', 12, f'{CL_MES}{HR + 1}:{CL_MES}{fim}')
dv(ws, 'F', len(ARGUMENTOS), f'{CL_ARG}{HR + 1}:{CL_ARG}{fim}')
dv(ws, 'G', len(CENAS), f'{CL_CENA}{HR + 1}:{CL_CENA}{fim}')
# Persona aceita mais de um valor (separados por ";"), então não leva lista suspensa.

# ---------- Visão do mês ----------
wv = wb.create_sheet('Visão do mês', 1)
wv['A1'] = 'Visão do mês: o capítulo e o equilíbrio'
wv['A1'].font = TITULO
wv['A2'] = 'Primeiro nível do planejamento. As contagens são automáticas e servem para a checagem de equilíbrio na revisão mensal.'
wv['A2'].font = Font(name=F, italic=True, size=9)
R = "'Calendário'"
rng = lambda c: f"{R}!${c}$5:${c}$400"
info = [
    ('Mês', 'Mês 1'),
    ('Período', f'=TEXT({R}!$D$2,"DD/MM")&" a "&TEXT({R}!$D$2+27,"DD/MM")'),
    ('Capítulo', '"Tudo começa com uma pergunta": quem somos, por que cuidar antes e como o cuidado funciona.'),
    ('Argumento principal', f'{LBE}, com a apresentação da marca. A NR-1 entra como contexto (artigo da RH Pra Você) e a ciência sustenta o mês (OMS, Diabetes Prevention Program, RAND).'),
    ('Mensagem do mês', 'Todo afastamento tem um antes. A Fair Health cuida nesse antes: dentro da empresa, com gente de verdade, começando por uma pergunta.'),
    ('Frase-tese', 'Não é benefício. É infraestrutura de cuidado.'),
    ('Restrições', 'Sem apresentar o time (questão administrativa), sem filmar clientes e sem nomes reais (personagens fictícias quando preciso). Vídeos de tipografia, motion, gravação de tela ou imagem de banco.'),
    ('Números', 'Sim, com suavidade: cuidado primeiro, número como contexto, fonte visível. Um número principal por peça.'),
    ('Datas do calendário', '02/11 feriado (Finados): a semana 3 começa na terça. 14/11 (sábado) Dia Mundial do Diabetes: post de ciência na véspera (M1-16) e stories no dia. Novembro Azul: só se houver conexão real.'),
]
for i, (k_, v_) in enumerate(info, 4):
    a = wv.cell(row=i, column=1, value=k_)
    b = wv.cell(row=i, column=2, value=v_)
    a.font, a.fill, a.border, a.alignment = CAB, FILL_CAB, BORDA, WRAP
    b.font, b.border, b.alignment = NORMAL, BORDA, WRAP
    wv.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
r0 = 4 + len(info) + 1
wv.cell(row=r0, column=1, value='A história semana a semana').font = Font(name=F, bold=True, size=12, color=TEXTO)
cab = ['Semana', 'Cena', 'Pergunta que responde', 'Peças (IDs)', 'Total', 'Vídeos', 'Institucional', 'Produto', 'Projetos']
SEMANAS = [
    (1, 'Abertura e Dor', '"Isso é comigo?" A marca se apresenta e o RH se reconhece no problema.', 'M1-01 a M1-04'),
    (2, 'Virada', '"Como a Fair Health resolve?" Método, tese e o que o RH recebe.', 'M1-05 a M1-08'),
    (3, 'Prova', '"Por que acreditar?" Adesão, acompanhamento, origem e autoridade.', 'M1-09 a M1-12'),
    (4, 'Convite', '"Como começo?" Para o colaborador, o líder e o RH, com a ciência por trás do cuidado.', 'M1-13 a M1-16'),
]
for j, t in enumerate(cab, 1):
    c = wv.cell(row=r0 + 1, column=j, value=t)
    c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
    c.alignment = Alignment(wrap_text=True, vertical='center')
for i, (s, cena, perg, ids) in enumerate(SEMANAS, r0 + 2):
    vals = [s, cena, perg, ids,
            f'=COUNTIFS({rng("C")},$A{i})',
            f'=COUNTIFS({rng("C")},$A{i},{rng(CL_FMT)},"{V}")',
            f'=COUNTIFS({rng("C")},$A{i},{rng(CL_LINHA)},"{I}")',
            f'=COUNTIFS({rng("C")},$A{i},{rng(CL_LINHA)},"{P}")',
            f'=COUNTIFS({rng("C")},$A{i},{rng(CL_LINHA)},"{J}")']
    for j, v_ in enumerate(vals, 1):
        c = wv.cell(row=i, column=j, value=v_)
        c.font, c.alignment, c.border = NORMAL, WRAP, BORDA
tot = r0 + 2 + len(SEMANAS)
wv.cell(row=tot, column=1, value='Total').font = CAB
for j in range(5, 10):
    L = get_column_letter(j)
    c = wv.cell(row=tot, column=j, value=f'=SUM({L}{r0 + 2}:{L}{tot - 1})')
    c.font, c.border = CAB, BORDA

r1 = tot + 2
wv.cell(row=r1, column=1, value='Equilíbrio por argumento e por cena').font = Font(name=F, bold=True, size=12, color=TEXTO)
for j, t in enumerate(['Argumento', 'Peças', '', 'Cena', 'Peças', '', 'Pautas com validação médica'], 1):
    if t:
        c = wv.cell(row=r1 + 1, column=j, value=t)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
for i, a_ in enumerate(ARGUMENTOS, r1 + 2):
    wv.cell(row=i, column=1, value=a_).font = NORMAL
    wv.cell(row=i, column=2, value=f'=COUNTIFS({rng(CL_ARG)},$A{i})').font = NORMAL
for i, c_ in enumerate(CENAS, r1 + 2):
    wv.cell(row=i, column=4, value=c_).font = NORMAL
    wv.cell(row=i, column=5, value=f'=COUNTIFS({rng(CL_CENA)},$D{i})').font = NORMAL
wv.cell(row=r1 + 2, column=7, value=f'=SUMPRODUCT(({rng("B")}<>"")*ISNUMBER(SEARCH("responsável médica",{rng(CL_VAL)})))').font = NORMAL
for j, w in enumerate([22, 16, 44, 16, 9, 9, 13, 10, 10], 1):
    wv.column_dimensions[get_column_letter(j)].width = w

# ---------- tabelas auxiliares ----------
def tabela(ws_, r0, titulo, cab_, dados):
    ws_.cell(row=r0, column=1, value=titulo).font = Font(name=F, bold=True, size=12, color=TEXTO)
    for j, t in enumerate(cab_, 1):
        c = ws_.cell(row=r0 + 1, column=j, value=t)
        c.font, c.fill, c.border = CAB, FILL_CAB, BORDA
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, linha in enumerate(dados, r0 + 2):
        for j, v in enumerate(linha, 1):
            c = ws_.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = NORMAL, WRAP, BORDA
    return r0 + 2 + len(dados) + 1

def larguras(ws_, ls):
    for j, w in enumerate(ls, 1):
        ws_.column_dimensions[get_column_letter(j)].width = w

# ---------- Formatos ----------
wf = wb.create_sheet('Formatos', 3)
larguras(wf, [26, 44, 44, 44, 40, 18])
r = tabela(wf, 1, 'Formatos do mês: sem time em cena, sem gravar nas empresas-clientes e sem nomes reais',
       ['Formato', 'Como é', 'Por que funciona', 'O que precisa para produzir', 'Cuidados', 'Onde usamos'], [
    ('Reel de tipografia animada', 'Frases curtas na tela, uma por vez, sobre o fundo da marca. Trilha calma; voz em off opcional.',
     'Gancho na primeira frase, lido sem som. Não depende de filmagem. Sequências do tipo "antes do..." prendem até o fim.',
     'Texto aprovado, Canva (animação de texto) ou CapCut. 1 a 2 horas.', 'Frases curtas; no máximo 7 cenas; logo só no fim.', 'M1-03, M1-11'),
    ('Explicador animado (motion com ícones)', 'Conceito explicado com ícones e voz em off, em até 60 s.',
     'Explica um conceito novo (MEV, passos da oferta) mais rápido que um carrossel e circula bem no LinkedIn.',
     'Roteiro aprovado, ícones lineares, voz em off (pode ser gravada pelo celular em local silencioso).', 'Validação médica de tudo o que for clínico.', 'M1-05, M1-15'),
    ('Gravação de tela', 'Rolagem por um painel ou tela do produto, com texto na tela.',
     'Mostra o produto de verdade sem expor ninguém e responde à dúvida "o que eu recebo?".',
     'Painel com dados fictícios e aviso "dados ilustrativos".', 'Nunca dado real nem individual.', 'M1-07'),
    ('Número contado como história', 'Primeiro a cena de cuidado, depois o número, em ritmo lento e com a fonte discreta.',
     'O número prova sem tirar o foco do cuidado; o RH leva a prova para a diretoria sem cara de anúncio.',
     'Números liberados (aba Fontes) ou de estudos (aba Ciência).', 'Sempre agregado; sem nome da empresa; regras de número suave (abaixo).', 'M1-09'),
    ('Carrossel "o que a ciência diz"', 'Um estudo por post: o que foi feito, o que encontrou, o que isso tem a ver com o nosso cuidado, e a referência completa.',
     'Transfere a autoridade da ciência para a marca e educa sem prometer. Conecta datas do calendário ao trabalho sem oportunismo.',
     'Estudo da aba Ciência, lido pela responsável médica.', 'Explicar para quem o estudo vale; nunca transformar em promessa individual.', 'M1-14, M1-16'),
    ('Reel de dicas com imagem de banco', '"3 hábitos" com texto na tela sobre imagens de banco.',
     'Dica prática é salva e compartilhada; imagens de banco evitam gravar pessoas.',
     'Banco de imagens (Canva, Pexels) sem rosto identificável.', 'Sem promessa clínica; validação médica.', 'M1-13'),
    ('Carrossel passo a passo', 'Um passo ou uma objeção por slide, com capa-pergunta.',
     'É o formato mais salvo em conteúdo explicativo e o que já testamos no carrossel de acolhimento.',
     'Modelos M1, M2 e M3 no Canva.', 'Até 10 slides; uma ideia por slide.', 'M1-02, M1-10'),
    ('História com personagem fictícia', 'Uma personagem inventada (nome, idade, rotina) mostra como o cuidado funciona na prática.',
     'História prende mais que explicação e permite mostrar o plano sem expor ninguém.',
     'Personagem escrita com a responsável médica para ser plausível.', 'Sinalizar "história ilustrativa, personagem fictícia" no post; nunca apresentar como depoimento.', 'M1-08'),
    ('Carrossel-resumo de artigo', 'As ideias principais de um artigo publicado, com citações literais e link nos stories.',
     'Leva a autoridade do veículo para a marca; muito compartilhado entre profissionais de RH.',
     'Artigo publicado.', 'Citar só o artigo e o veículo, sem nome do autor; citações literais conferidas.', 'M1-04'),
    ('Carrossel comparativo (isto × aquilo)', 'Duas colunas que contrastam duas ideias.',
     'A diferença fica clara em segundos; ajuda o RH a defender a decisão internamente.',
     'Modelo M1 e comparativo em duas colunas.', 'Não citar marcas de terceiros.', 'M1-06'),
    ('Card de citação', 'Uma frase em destaque com assinatura.',
     'Fácil de compartilhar nos stories e no LinkedIn; reaproveita conteúdo sem repetir.',
     'Modelo M8.', 'Citação literal; crédito ao artigo, sem nome de pessoa.', 'M1-12'),
    ('Checklist salvável', 'Perguntas ou itens numerados, um por slide, com convite para salvar.',
     'Conteúdo útil é salvo e repassado à liderança.', 'Modelo M2 com etiqueta 1/5.', 'Fonte oficial quando cita norma.', 'Banco de ideias (NR-1)'),
])
tabela(wf, r, 'Números com suavidade (decisão de 09/10: números sim, foco no cuidado)', ['Regra', 'Como fazer', 'Exemplo', 'Evitar', '', ''], [
    ('Cuidado primeiro', 'A frase de cuidado abre; o número entra depois, como contexto.', '"Ninguém volta a um lugar onde não se sentiu cuidado." → "8 em cada 10 voltaram."', 'Abrir com o número como manchete.', '', ''),
    ('Um número principal por peça', 'Os outros números, se houver, vão em texto menor ou na legenda.', 'M1-09: 8 em cada 10 em destaque; 98% em texto menor.', 'Painel de números em sequência.', '', ''),
    ('Traduzir em gente', 'Preferir "8 em cada 10 pessoas" a "81,3%"; arredondar com honestidade.', '"Mais de 470 mil licenças" em vez de "472.328".', 'Casas decimais e jargão estatístico.', '', ''),
    ('Tom calmo', 'Verde-escuro, ritmo lento, sem efeito de conquista.', 'Número em Raleway, peso médio, sem seta.', 'Vermelho, alarme, "explodiu", "recorde assustador".', '', ''),
    ('Fonte sempre', 'Fonte visível no rodapé e referência completa na aba Ciência ou Fontes.', '"Fonte: OMS, 2022."', 'Número sem fonte ou "estudos mostram".', '', ''),
    ('Sem promessa', 'O número fala do estudo ou do grupo, nunca do resultado de quem vai ser atendido.', '"No estudo, 58% menos casos novos."', '"Você vai reduzir 58% do seu risco."', '', ''),
])

# ---------- Ciência ----------
wc = wb.create_sheet('Ciência', 4)
larguras(wc, [30, 52, 46, 40, 40, 34, 16])
tabela(wc, 1, 'Ciência: estudos e fontes oficiais para posts (sempre com a referência completa)',
       ['Estudo ou fonte', 'Referência completa', 'O que encontrou', 'Como apresentar com suavidade', 'Ligação com o nosso trabalho', 'Cuidados', 'Onde usamos'], [
    ('Diabetes Prevention Program (2002)', DPP,
     '3.234 adultos com pré-diabetes, acompanhados por 2,8 anos em média: mudança de estilo de vida (alimentação, 7% de perda de peso, 150 min de atividade física por semana) reduziu em 58% os casos novos de diabetes tipo 2 em relação ao placebo; a metformina reduziu 31%.',
     '"Quem mudou a rotina teve 58% menos casos novos." O 31% do medicamento em texto menor.',
     'A base da Medicina do Estilo de Vida: o cuidado começa pela rotina.', 'Vale para adultos com risco aumentado; não é promessa individual. Validação médica.', 'M1-16'),
    ('OMS: diretrizes sobre saúde mental no trabalho (2022)', OMS,
     'Primeiras diretrizes da OMS sobre o tema. Recomenda, entre outras ações, preparar gestores para prevenir ambientes de estresse e acolher quem está em sofrimento. Com a OIT, estima que depressão e ansiedade custam cerca de 12 bilhões de dias de trabalho por ano no mundo.',
     'Recomendação como gesto simples para o líder; os 12 bilhões de dias, se usados, em texto menor.',
     'Liderança Bem-Estar e NR-1: preparar líderes e cuidar dentro da empresa.', 'Conferir a redação exata nas diretrizes.', 'M1-14'),
    ('Licenças por transtornos mentais no Brasil (2024)', INSS,
     '472.328 licenças por transtornos mentais em 2024, o maior número da série iniciada em 2014. Ansiedade e depressão lideram.',
     '"Mais de 470 mil licenças. Cada uma teve um antes." Em texto calmo, depois da frase de cuidado.',
     'O eixo "cuidar antes".', 'Os números divulgados variam entre veículos (440 mil e 472 mil): conferir no painel oficial e citar a fonte primária.', 'M1-03'),
    ('Why sleep matters (RAND Europe, 2016)', RAND,
     'Sono insuficiente custa até 2,28% do PIB dos EUA e cerca de 1,2 milhão de dias de trabalho por ano; dormir menos de 6 horas está associado a risco de mortalidade 13% maior do que dormir de 7 a 9 horas.',
     'Usar a faixa saudável (7 a 9 horas) para o colaborador; o custo em dias de trabalho só para o RH, na legenda do LinkedIn.',
     'Pilar sono.', 'Estimativas modeladas ("até"); evitar o dado de mortalidade em post para colaborador.', 'M1-13'),
    ('Global Burden of Disease 2019 (citado no artigo da RH Pra Você)', 'GBD 2019 Mental Disorders Collaborators; artigo na Revista Brasileira de Epidemiologia (SciELO): https://www.scielo.br/j/rbepid/a/cwDkPDjSkDcC3hB73kp7yfC/',
     'Pessoas vivendo com 12 transtornos mentais em 204 países: de 654,8 milhões (1990) para 970,1 milhões (2019), +48,1%.',
     '"Quase 1 bilhão de pessoas", em texto menor, como contexto.', 'Saúde mental é multifatorial; a empresa cuida da parte dela.', 'Citar a fonte primária.', 'M1-04'),
    ('Atividade física e saúde mental (Singh et al., 2023)', 'Singh B. et al. Effectiveness of physical activity interventions for improving depression, anxiety and distress: an overview of systematic reviews. Br J Sports Med. 2023. doi:10.1136/bjsports-2022-106195',
     'Revisão de 97 revisões (1.039 estudos, mais de 128 mil participantes): atividade física teve efeito moderado na redução de sintomas de depressão, ansiedade e sofrimento psicológico.',
     '"Mexer o corpo também cuida da mente", com o efeito descrito em palavras, sem número.',
     'Pilar atividade física e manejo do estresse.', 'Não usar o "1,5 vez melhor que remédio" da imprensa: não está no resumo do artigo.', 'Banco de ideias'),
    ('Relações sociais e mortalidade (Holt-Lunstad et al., 2010)', 'Holt-Lunstad J., Smith T.B., Layton J.B. Social relationships and mortality risk: a meta-analytic review. PLoS Med. 2010;7(7):e1000316. doi:10.1371/journal.pmed.1000316',
     '148 estudos, 308.849 participantes: relações sociais mais fortes associadas a maior chance de sobrevivência (razão de chances 1,50), efeito comparável a fatores de risco conhecidos.',
     '"Vínculos também são saúde", sem o número na peça para colaborador.',
     'Pilar relacionamentos; ambiente de trabalho.', 'Evitar o "equivale a 15 cigarros por dia" da imprensa: não está no artigo.', 'Banco de ideias'),
    ('Estressores no trabalho, mortalidade e custos (Goh, Pfeffer e Zenios, 2016)', 'Goh J., Pfeffer J., Zenios S.A. The relationship between workplace stressors and mortality and health costs in the United States. Management Science. 2016;62(2):608-628. doi:10.1287/mnsc.2014.2115',
     'Modelo com 10 estressores do trabalho (jornadas longas, insegurança, baixa autonomia, conflito trabalho-família, entre outros) estima mais de 120 mil mortes por ano e de 5% a 8% dos gastos com saúde nos EUA.',
     'Para o RH: "o jeito de trabalhar também entra na conta da saúde", número em texto menor.',
     'NR-1 e riscos psicossociais; Stop Loss.', 'Conferir os números no artigo original antes de usar (vieram de resumos).', 'Banco de ideias'),
])

# ---------- Fontes ----------
wfo = wb.create_sheet('Fontes', 5)
larguras(wfo, [30, 60, 40, 40, 18])
tabela(wfo, 1, 'Fontes internas, o artigo e o que não usar (nomes de clientes só para uso interno, nunca no post)', ['Fonte', 'O que diz', 'Como usar', 'Situação', 'Onde usamos'], [
    ('Artigo "Riscos psicossociais no ambiente de trabalho e saúde" (RH Pra Você, 25/09/2026)',
     'Desde maio de 2026 a NR-1 exige que as empresas identifiquem, avaliem e gerenciem os riscos psicossociais; o trabalho pode adoecer, mas não é o único vilão; a empresa responde pelo ambiente que constrói, mas não controla todas as variáveis; responsabilidade compartilhada (poder público, sociedade, empresas e indivíduo); "Saúde mental é, ao mesmo tempo, de cada um e de todos." URL: https://rhpravoce.com.br/colab/riscos-psicossociais-no-ambiente-de-trabalho-2',
     'Resumo em carrossel, card de citação, link nos stories. Citar só o artigo e o veículo, sem nome do autor. Citações sempre literais.', 'Publicado.', 'M1-04, M1-12'),
    ('Cliente do setor industrial, 1º semestre de 2026 (uso interno: Chemitec)', '161 atendimentos, 48 colaboradores atendidos (98% de adesão), 81,3% voltaram para 2 ou mais atendimentos, 46 smartbands em uso.',
     'Sempre agregado e como "uma indústria". Nunca o nome da empresa.', 'Liberado (instagram/voice.md).', 'M1-09'),
    ('Hospital onde o modelo começou, 2025 (uso interno: Hospital Dia Santo Amaro)', 'Reclamações mensais de 34 para cerca de 20 (cerca de 40% a menos); conformidade média de 98,2%; apresentado em congresso internacional.',
     'História de origem e prova do modelo, sem o nome do hospital.', 'Liberado (instagram/voice.md).', 'Banco de ideias'),
    ('Harvard Business Review (2022) e SHRM (2023)', '−28% de absenteísmo (HBR 2022); −22% de custos com saúde (SHRM 2023).',
     'Argumento de retorno, sempre com a fonte. Fica para o mês de Stop Loss.', 'Liberado (instagram/voice.md).', 'Banco de ideias'),
    ('ACLM e CBMEV', 'Definição de Medicina do Estilo de Vida e seus 6 pilares: alimentação, atividade física, sono, manejo do estresse, relacionamentos e evitar substâncias de risco.', 'Base do explicador de MEV.', 'Responsável médica confere. Espiritualidade fora (a confirmar).', 'M1-05'),
    ('Não usar', '"23% mais caro", "R$ 2.400 por colaborador", "maior índice de burnout" (sem fonte no posicionamento); −27%, "R$ 2,70 por R$ 1" e "70% dos custos do plano"; "validado no SUS" até confirmar.',
     '—', 'Bloqueado até ter fonte.', '—'),
])

# ---------- Banco de ideias ----------
wb_ = wb.create_sheet('Banco de ideias', 6)
larguras(wb_, [40, 26, 18, 52, 30])
tabela(wb_, 1, 'Banco de ideias: o que ficou para depois e os próximos meses',
       ['Ideia', 'Origem', 'Status', 'Motivo / o que destrava', 'Próximo uso sugerido'], [
    ('Série Quem cuida (apresentação de cada integrante do time)', 'Calendário de 90 dias', 'Adiada',
     'Não vamos apresentar o time por agora (questão administrativa). As capas já feitas ficam guardadas em instagram/posts/time.', 'Quando a questão for resolvida'),
    ('Um dia no Escritório do Cuidado (vídeo no espaço)', 'Calendário de 90 dias (M1-03)', 'Adiada',
     'Não podemos filmar clientes: o espaço fica dentro das empresas-clientes.', 'Recriar em local próprio ou em motion'),
    ('Bastidores de uma implantação', 'Calendário de 90 dias (M2-01)', 'Adiada', 'Mesmo motivo: gravação dentro da empresa-cliente.', 'Versão em motion, com empresa fictícia'),
    ('Depoimentos em vídeo e por escrito', 'Calendário de 90 dias (M2-10, M3-03)', 'Adiada',
     'Não podemos filmar clientes; nunca nomes; limites do CFM para depoimentos.', 'Trocar por histórias com personagem fictícia'),
    ('Workshops nos pilares (fotos reais)', 'Calendário de 90 dias (M1-06, M2-13)', 'Aguardando material',
     'Fotos de encontros expõem clientes. Possível versão só com materiais e slides do workshop.', 'Mês de Ambiente positivo'),
    ('Registro do mês', 'Calendário de 90 dias', 'Aguardando pauta', 'Depende da agenda de eventos e entregas.', 'Quando houver evento próprio'),
    ('Checklist "Antes de preencher a planilha: 5 perguntas sobre riscos psicossociais"', 'Mês 1 (versão de 09/10)', 'Ideia',
     'Saiu do Mês 1 para dar lugar ao post de ciência do Dia Mundial do Diabetes. Abre bem o mês da NR-1.', 'Mês 2, semana 1'),
    ('Série NR-1 (contexto, diagnóstico, cuidado individualizado, acompanhamento, versão para colaboradores)', 'Plano mensal (Mês 2)', 'Ideia',
     'Próximo capítulo. O artigo da RH Pra Você já abre o tema. Estudo de Goh, Pfeffer e Zenios (aba Ciência) para o RH.', 'Mês 2'),
    ('Série "Ciência no cuidado": um estudo por mês ligado ao pilar ou ao argumento do mês', 'Decisão de 09/10', 'Ideia',
     'Estudos na aba Ciência: atividade física e saúde mental (Singh, 2023), relações sociais (Holt-Lunstad, 2010), estressores do trabalho (Goh, 2016), OMS (12 bilhões de dias).', 'Todo mês'),
    ('"O que nos diferencia": plano de saúde, plataformas de bem-estar e QVT (3 posts)', 'Posicionamento v3.2', 'Ideia', 'Falar da categoria, sem citar marcas.', 'Mês 2 ou 3'),
    ('"Quem cuida também precisa de cuidado" (edição especial para a saúde)', 'Posicionamento v3.2 (arenas)', 'Ideia', 'Liga a origem hospitalar ao cuidado de quem cuida.', 'Mês 3'),
    ('Pilares da MEV, um por vez (alimentação, atividade física, manejo do estresse, relacionamentos, evitar substâncias de risco)', 'Calendário de 90 dias', 'Ideia',
     '6 pilares, sem espiritualidade (a confirmar). Sono já entra no Mês 1.', 'Um pilar por mês'),
    ('Histórias com personagens fictícias (um dia na rotina de quem é cuidado)', 'Decisão de 09/10', 'Ideia', 'Sempre sinalizadas como ilustrativas.', 'Todo mês, 1 peça'),
    ('Case do hospital onde o modelo começou (sem o nome)', 'voice.md (liberado)', 'Ideia', 'Prova da origem do modelo.', 'Mês de Stop Loss ou de prova'),
    ('Indicadores para o RH (HBR 2022, SHRM 2023)', 'voice.md (liberado)', 'Ideia', 'Argumento de retorno, um indicador por vez.', 'Mês de Stop Loss'),
])

# ---------- Linhas e personas ----------
wp = wb.create_sheet('Linhas e personas', 7)
larguras(wp, [24, 46, 46, 40, 34])
r = tabela(wp, 1, 'Linhas editoriais', ['Linha', 'O que entra', 'Objetivos típicos', 'Formatos que funcionam agora', 'Peso no mês'], [
    (I, 'O que é a Fair Health, por que existe, de onde veio (Escritório do Paciente → Escritório do Cuidado, ecossistema FairJob) e em que acredita ("cuidar antes", "não é benefício, é infraestrutura de cuidado").',
     'Apresentar, gerar confiança, fixar a tese.', 'Tipografia animada, linha do tempo, comparativo, card de pergunta.', 'Cerca de 1/3'),
    (P, 'Escritório do Cuidado (percurso, sigilo), Medicina do Estilo de Vida e os pilares, plano individualizado, tecnologia e temas do entorno (NR-1, saúde mental) quando a ligação com a oferta é clara.',
     'Explicar, educar com rigor, tirar dúvidas, levar à conversa comercial.', 'Carrossel passo a passo, explicador animado, checklist, dicas com imagem de banco.', 'Cerca de 1/3 a 2/5'),
    (J, 'O que a Fair Health está fazendo e entregando: números liberados, entregas para o RH, participações na imprensa e em eventos, como chega a uma empresa.',
     'Provar, mostrar movimento, gerar autoridade.', 'Dado animado, gravação de tela, resumo de artigo, card de citação.', 'Cerca de 1/4 a 1/3'),
])
tabela(wp, r, 'Personas (posicionamento v3.2)', ['Persona', 'Quem é', 'Dor de entrada', 'Como falar', 'Cuidados'], [
    (RH, 'CHROs, gerentes de pessoas e business partners de empresas de médio e grande porte, principalmente indústria. Público principal.',
     'Como atender à NR-1 de um jeito que mude algo? Meu programa tem adesão? O que eu recebo? Como começar?', 'Assertivo e pragmático, com processo e evidência.', 'Números só liberados e com fonte.'),
    (LID, 'Gestores e líderes de equipe.', 'Como cuidar do time no dia a dia sem ser médico nem psicólogo?', 'Direto, com perguntas práticas.', 'Não transferir ao líder a responsabilidade clínica.'),
    (COL, 'Quem é ou pode ser atendido pelo Escritório do Cuidado.', 'Vou ser julgado? O que eu contar fica entre nós? Cabe na minha rotina?', 'Acolhedor, em segunda pessoa, sem números de negócio.', 'Nunca expor quem é atendido.'),
    (MT, 'Médicos do trabalho e equipes de SST.', 'Isso substitui ou complementa o que já fazemos?', 'Técnico, com fontes e precisão.', 'Validação da responsável médica.'),
    (RC, 'Áreas de risco, compliance e jurídico.', 'Como documentar e acompanhar os riscos psicossociais?', 'Objetivo, com a norma e a fonte.', 'Não prometer conformidade.'),
    (CFO, 'Diretoria financeira.', 'Qual o retorno? O que muda na sinistralidade?', 'Indicadores com fonte, sem promessa.', 'Argumento de retorno no mês de Stop Loss.'),
    (CL, 'C-Level e relações com investidores (ESG).', 'Como o cuidado aparece nos indicadores da empresa?', 'Estratégico, curto.', 'Sem números sem fonte.'),
])

# ---------- Pendências ----------
wpd = wb.create_sheet('Pendências', 8)
larguras(wpd, [5, 60, 30, 34, 30])
pend = [
    ('Questão administrativa que adia a apresentação do time: previsão para retomar a série Quem cuida.', 'Sócios', 'Banco de ideias', 'Aberta'),
    ('Onde gravar o vídeo de abertura sem time e sem empresa-cliente (local próprio ou neutro) ou montar com fotos e tipografia.', 'Comunicação', 'M1-01', 'Aberta'),
    ('Assistentes sociais fazem parte do atendimento? (o posicionamento cita médico e psicológico)', 'Time Fair Health', 'M1-01', 'Aberta'),
    ('Nome da responsável médica que valida o conteúdo clínico e lê os estudos da aba Ciência.', 'Time Fair Health', 'Todas as pautas com "responsável médica"; M1-13, M1-14, M1-16', 'Aberta'),
    ('Pilares da MEV: 6 (ACLM), sem espiritualidade. Decisão de 09/10, ainda sem certeza.', 'Responsável médica + sócios', 'M1-05 e série Pilares', 'Aberta'),
    ('Conferir no painel oficial do Ministério da Previdência o número de licenças por transtornos mentais em 2024 (veículos divulgam 440 mil e 472 mil).', 'Comunicação', 'M1-03', 'Aberta'),
    ('Nome único do plano: Plano Individualizado de Cuidado ou Plano de Cuidado Personalizado (PPC).', 'Sócios', 'M1-08', 'Aberta'),
    ('Etapas e prazos clínicos: bioimpedância, smartband por 48 h, retorno.', 'Responsável médica', 'M1-02, M1-08, M1-10', 'Aberta'),
    ('Fonte de "validado no SUS" e datas da história de origem.', 'Time Fair Health', 'M1-11', 'Aberta'),
    ('Indicadores que entram no relatório do RH (fairdata).', 'Time Fair Health + fairdata', 'M1-07', 'Aberta'),
    ('Collab no Instagram e LinkedIn pela FairJob: quem publica, quem aceita o convite.', 'Comunicação + FairJob', 'Todas as pautas', 'Aberta'),
    ('Destino do link e da chamada CUIDADO (www.fairhealth.com.br) e quem responde comentários e DMs.', 'Comunicação', 'M1-06, M1-15', 'Aberta'),
    ('Texto, prazos e obrigações da NR-1 com fonte oficial; aprovação da posição [SUGESTÃO] 1 do voice.md.', 'Responsável médica + sócios', 'Banco de ideias (checklist NR-1)', 'Aberta'),
    ('Fontes dos números do posicionamento (23%, R$ 2.400, burnout).', 'Sócios', 'Fontes', 'Aberta'),
    ('Crédito do autor do artigo da RH Pra Você.', 'Comunicação', 'M1-04, M1-12', 'Resolvida em 09/10: citar só o artigo e o veículo, sem nome.'),
    ('Números no Mês 1 e nome do cliente industrial.', 'Sócios', 'M1-09', 'Resolvida em 09/10: números sim, com suavidade; nunca o nome do cliente.'),
]
tabela(wpd, 1, 'Pendências e validações', ['#', 'Pendência', 'Quem confirma', 'Pautas afetadas', 'Situação'],
       [(i, a, b, c, d) for i, (a, b, c, d) in enumerate(pend, 1)])

wb.move_sheet('Listas', offset=len(wb.sheetnames))
wb.save(OUT)
print('ok', OUT)
